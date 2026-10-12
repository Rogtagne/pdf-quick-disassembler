"""
PDF Quick Disassembler  v3.6  —  Tablas Fieles
Produced by GoTopo  —  2026

FASE 1: DOCX / PPTX  →  PDF
FASE 2: PDF  →  PPTX editable  (fondo fiel o modo objetos independientes)
        PDF  →  DOCX editable   (texto corrido que Word reacomoda)
FASE 3: PDF con tablas de Excel  →  XLSX + CSV
"""

import os
import sys
import threading
import subprocess
import tempfile
import shutil
from pathlib import Path

from utils.logging import get_logger

log = get_logger("topo_vs_pdf")

# Tkinter solo se carga en modo desktop (no en servidor web Linux)
_WEB_MODE = (os.environ.get("PQD_WEB_MODE") == "1"
             or any("gunicorn" in a for a in sys.argv))

if not _WEB_MODE:
    try:
        from tkinter import (
            Tk, ttk, StringVar, Label, Button, Frame,
            filedialog, messagebox
        )
        from tkinter.ttk import Progressbar, Notebook
        _TKINTER_OK = True
    except ImportError:
        _TKINTER_OK = False
else:
    _TKINTER_OK = False

# ── Poppler ───────────────────────────────────────────────────────────────────
def _setup_poppler():
    # Linux / Render.com: pdftoppm instalado via apt (poppler-utils)
    for lp in ["/usr/bin", "/usr/local/bin", "/opt/homebrew/bin"]:
        if Path(lp, "pdftoppm").exists():
            return lp  # ya en PATH del sistema, solo retornar la ruta

    # Windows: buscar en C:\Poppler y rutas comunes
    candidates = []
    poppler_root = Path(r"C:\Poppler")
    if poppler_root.exists():
        for sub in sorted(poppler_root.iterdir(), reverse=True):
            for rel in ["Library/bin", "bin"]:
                p = sub / rel
                if p.exists() and any(p.glob("pdftoppm*")):
                    candidates.insert(0, str(p))
    candidates += [
        r"C:\Program Files\poppler\Library\bin",
        r"C:\Program Files\poppler\bin",
    ]
    for c in candidates:
        if Path(c).exists():
            if c not in os.environ.get("PATH", ""):
                os.environ["PATH"] = c + os.pathsep + os.environ.get("PATH", "")
            return c
    return None

POPPLER_PATH = _setup_poppler()
if POPPLER_PATH:
    log.debug("Poppler encontrado en %s", POPPLER_PATH)
else:
    log.warning("Poppler no encontrado: el renderizado de páginas (pdf2image) fallará")

# ── Paleta de colores ─────────────────────────────────────────────────────────
C_BG      = "#12172B"   # fondo profundo
C_PANEL   = "#1C2240"   # paneles
C_CARD    = "#252D50"   # tarjetas
C_CARD2   = "#1E2840"   # tarjeta secundaria
C_ACCENT  = "#4F7EFF"   # azul principal
C_GREEN   = "#00C896"   # verde éxito
C_WARN    = "#FF6B6B"   # error
C_GOLD    = "#F5C518"   # dorado GoTopo
C_TEXT    = "#E8ECFF"   # texto principal
C_MUTED   = "#7880AA"   # texto secundario
C_BORDER  = "#2E3860"   # bordes


# ══════════════════════════════════════════════════════════════════════════════
#  CONVERSIÓN — FASE 1
# ══════════════════════════════════════════════════════════════════════════════

def _libre_convert(src: Path, dst_dir: Path) -> Path:
    for soffice in ["soffice",
                    r"C:\Program Files\LibreOffice\program\soffice.exe",
                    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]:
        try:
            r = subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf",
                 "--outdir", str(dst_dir), str(src)],
                capture_output=True, timeout=120)
            out = dst_dir / (src.stem + ".pdf")
            if r.returncode == 0 and out.exists():
                log.debug("LibreOffice (%s) convirtió %s", soffice, src.name)
                return out
            log.debug("LibreOffice (%s) falló: código=%s stderr=%s", soffice,
                      r.returncode, r.stderr.decode(errors="replace").strip()[-300:])
        except (FileNotFoundError, subprocess.TimeoutExpired) as e:
            log.debug("LibreOffice (%s) no disponible: %s", soffice, type(e).__name__)
            continue
    raise RuntimeError("LibreOffice no encontrado")


def _docx_to_pdf_python(src: Path, dst: Path):
    from docx import Document
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    doc  = Document(str(src))
    pdfd = SimpleDocTemplate(str(dst), pagesize=letter,
                             leftMargin=2*cm, rightMargin=2*cm,
                             topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    story  = []
    for para in doc.paragraphs:
        text  = para.text.strip()
        pname = (para.style.name or "") if para.style else ""
        if not text:
            story.append(Spacer(1, 6)); continue
        st = (styles["h1"] if pname.startswith("Heading 1") else
              styles["h2"] if pname.startswith("Heading 2") else
              styles["Normal"])
        story.append(Paragraph(text, st))
        story.append(Spacer(1, 4))
    if not story:
        story.append(Paragraph("(documento vacío)", styles["Normal"]))
    pdfd.build(story)


def _pptx_to_pdf_python(src: Path, dst: Path):
    from pptx import Presentation as _Prs
    from reportlab.pdfgen import canvas as rl_canvas
    prs  = _Prs(str(src))
    w_pt = prs.slide_width  / 12700   # type: ignore
    h_pt = prs.slide_height / 12700   # type: ignore
    c    = rl_canvas.Canvas(str(dst), pagesize=(w_pt, h_pt))
    for i, slide in enumerate(prs.slides):
        c.setPageSize((w_pt, h_pt))
        c.setFont("Helvetica-Bold", 14)
        c.setFillColorRGB(0.2, 0.2, 0.6)
        c.drawCentredString(w_pt/2, h_pt/2 + 16, f"Diapositiva {i+1}")
        c.setFont("Helvetica", 8)
        c.setFillColorRGB(0.5, 0.5, 0.5)
        c.drawCentredString(w_pt/2, h_pt/2 - 8,
                            "(instale LibreOffice para fidelidad total)")
        y = h_pt - 36
        for shape in slide.shapes:
            if not shape.has_text_frame: continue
            for para in shape.text_frame.paragraphs:  # type: ignore
                line = para.text.strip()
                if line and y > 16:
                    c.setFont("Helvetica", 8)
                    c.setFillColorRGB(0.15, 0.15, 0.15)
                    c.drawCentredString(w_pt/2, y, line[:90])
                    y -= 11
        c.showPage()
    c.save()


def fase1_convert(src_path: str, dst_path: str, progress_cb=None) -> bool:
    src = Path(src_path)
    dst = Path(dst_path)
    ext = src.suffix.lower()
    if progress_cb: progress_cb(10, f"Analizando {src.name}…")
    tmp = Path(tempfile.mkdtemp())
    try:
        try:
            if progress_cb: progress_cb(30, "Convirtiendo con LibreOffice…")
            shutil.copy2(_libre_convert(src, tmp), dst)
            if progress_cb: progress_cb(100, "✅  Conversión completada con LibreOffice")
            return True
        except Exception as e:
            log.warning("Fase 1: LibreOffice no disponible (%s); usando motor Python "
                        "(solo texto) para %s", e, src.name)
        if progress_cb: progress_cb(50, "Motor Python (solo texto)…")
        if ext == ".docx":   _docx_to_pdf_python(src, dst)
        elif ext == ".pptx": _pptx_to_pdf_python(src, dst)
        else: raise ValueError(f"Formato no soportado: {ext}")
        if progress_cb: progress_cb(100, "✅  Completado (instale LibreOffice para imágenes)")
        return True
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ══════════════════════════════════════════════════════════════════════════════
#  CONVERSIÓN — FASE 2
# ══════════════════════════════════════════════════════════════════════════════

def _is_bullet_or_numbered(text: str) -> bool:
    """
    Detecta inicios típicos de listas para evitar que se peguen al párrafo
    anterior cuando se agrupa el texto extraído del PDF.
    """
    import re

    return bool(re.match(r"^(\u2022|\-|\*|\d+[\.)]|[a-zA-Z][\.)])\s+", text.strip()))


def _join_paragraph_line(previous: str, new_line: str) -> str:
    """
    Une líneas consecutivas en un párrafo continuo.

    Caso especial:
        "administra-" + "ción"  ->  "administración"
    """
    previous = previous.rstrip()
    new_line = new_line.strip()

    if not previous:
        return new_line

    if previous.endswith("-") and new_line:
        return previous[:-1] + new_line

    return previous + " " + new_line


def _extract_text_lines(pdf_path: str):
    """
    Extrae texto embebido del PDF y lo devuelve agrupado por líneas.

    Esta función conserva el comportamiento original del MVP, pero queda
    separada para poder ofrecer dos modos en la Fase 2:

        1) Líneas exactas: máxima fidelidad posicional, muchas cajas pequeñas.
        2) Párrafos editables: menos cajas, edición más natural.

    Retorno:
        [(page_width_pt, page_height_pt, [line_blocks]), ...]
    """
    import pdfplumber

    pages = []

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            pw, ph = float(page.width), float(page.height)

            words = page.extract_words(
                x_tolerance=3,
                y_tolerance=3,
                keep_blank_chars=False,
                use_text_flow=False,
                extra_attrs=["size", "fontname"],
            )

            if not words:
                pages.append((pw, ph, []))
                continue

            # Agrupar primero por coordenada vertical aproximada.
            raw_lines: dict = {}
            for word in words:
                key = round(float(word.get("top", 0)) / 4) * 4
                raw_lines.setdefault(key, []).append(word)

            line_blocks = []

            for _, line_words in sorted(raw_lines.items()):
                line_words = sorted(line_words, key=lambda item: float(item["x0"]))

                # Si una misma altura contiene columnas o textos muy separados,
                # se separan en segmentos para no mezclar columnas o números.
                segments = []
                current_segment = []
                previous_x1 = None
                gap_limit = max(42.0, pw * 0.18)

                for word in line_words:
                    x0 = float(word["x0"])
                    if (
                        previous_x1 is not None
                        and current_segment
                        and (x0 - previous_x1) > gap_limit
                    ):
                        segments.append(current_segment)
                        current_segment = []
                    current_segment.append(word)
                    previous_x1 = float(word["x1"])

                if current_segment:
                    segments.append(current_segment)

                for segment in segments:
                    text = " ".join(str(word["text"]) for word in segment).strip()
                    if not text:
                        continue

                    x0 = float(min(word["x0"] for word in segment))
                    y0 = float(min(word.get("top", 0) for word in segment))
                    x1 = float(max(word["x1"] for word in segment))
                    y1 = float(max(word.get("bottom", y0 + 12) for word in segment))

                    first = segment[0]
                    size = float(first.get("size", 11) or 11)
                    fontname = str(first.get("fontname", "") or "")
                    bold = "Bold" in fontname or "bold" in fontname

                    line_blocks.append(
                        {
                            "text": text,
                            "x0": x0,
                            "y0": y0,
                            "x1": x1,
                            "y1": y1,
                            "size": size,
                            "bold": bold,
                            "height": max(y1 - y0, size, 8),
                            "kind": "line",
                        }
                    )

            pages.append((pw, ph, line_blocks))

    return pages


def _group_lines_into_paragraphs(line_pages):
    """
    Reconstruye párrafos aproximados a partir de líneas de texto PDF.

    v3.3 mejora dos puntos importantes:
        1) Conserva la lista interna de líneas originales dentro de cada párrafo.
           Esto permite mantener saltos de línea visuales dentro de una sola caja
           editable de PowerPoint, evitando que el layout se desacomode.
        2) Calcula ancho y alto usando la caja real del párrafo en el PDF, no un
           tamaño genérico. Así la ventana del párrafo queda mucho más correlacionada
           con el ancho visual original.

    Importante:
        Un PDF no conserva párrafos reales como un DOCX. Normalmente solo conserva
        caracteres/palabras con coordenadas. Por eso esta función usa heurísticas
        controladas: separación vertical, sangría, bullets, cambios de tamaño y
        títulos en negrita.
    """
    paragraph_pages = []

    for pw, ph, lines in line_pages:
        if not lines:
            paragraph_pages.append((pw, ph, []))
            continue

        paragraphs = []
        current = None

        for line in sorted(lines, key=lambda item: (item["y0"], item["x0"])):
            if current is None:
                current = {
                    "text": line["text"],
                    "lines": [line["text"]],
                    "line_heights": [float(line.get("height", 8))],
                    "line_widths": [float(line["x1"] - line["x0"])],
                    "x0": line["x0"],
                    "y0": line["y0"],
                    "x1": line["x1"],
                    "y1": line["y1"],
                    "size": line["size"],
                    "bold": line["bold"],
                    "last_line": line,
                    "line_count": 1,
                }
                continue

            previous = current["last_line"]
            vertical_gap = float(line["y0"] - previous["y1"])
            avg_height = max(float(previous.get("height", 8)), float(line.get("height", 8)), 8)
            indent_change = abs(float(line["x0"] - previous["x0"]))
            size_change = abs(float(line["size"] - previous["size"]))
            same_baseline = abs(float(line["y0"] - previous["y0"])) < avg_height * 0.40

            new_paragraph = False

            # Si hay dos segmentos en la misma altura pero muy separados en X,
            # probablemente son columnas/campos distintos y no un mismo párrafo.
            if same_baseline and float(line["x0"]) > float(previous["x1"]) + 35:
                new_paragraph = True

            # Separación vertical claramente mayor al interlineado normal.
            if vertical_gap > avg_height * 0.85:
                new_paragraph = True

            # Cambio fuerte de sangría: típico de título, lista o nuevo bloque.
            if indent_change > 28 and vertical_gap > avg_height * 0.18:
                new_paragraph = True

            # Listas: cada bullet/numeral conviene conservarlo como bloque editable.
            if _is_bullet_or_numbered(line["text"]):
                new_paragraph = True

            # Título breve en negrita seguido de texto: separar título y cuerpo.
            if previous["bold"] and len(previous["text"]) < 90 and vertical_gap > avg_height * 0.20:
                new_paragraph = True

            # Cambio importante de tamaño de fuente: probable encabezado o nota.
            if size_change >= 2.0 and vertical_gap > avg_height * 0.10:
                new_paragraph = True

            if new_paragraph:
                paragraphs.append(current)
                current = {
                    "text": line["text"],
                    "lines": [line["text"]],
                    "line_heights": [float(line.get("height", 8))],
                    "line_widths": [float(line["x1"] - line["x0"])],
                    "x0": line["x0"],
                    "y0": line["y0"],
                    "x1": line["x1"],
                    "y1": line["y1"],
                    "size": line["size"],
                    "bold": line["bold"],
                    "last_line": line,
                    "line_count": 1,
                }
            else:
                current["text"] = _join_paragraph_line(current["text"], line["text"])
                current["lines"].append(line["text"])
                current["line_heights"].append(float(line.get("height", 8)))
                current["line_widths"].append(float(line["x1"] - line["x0"]))
                current["x0"] = min(current["x0"], line["x0"])
                current["y0"] = min(current["y0"], line["y0"])
                current["x1"] = max(current["x1"], line["x1"])
                current["y1"] = max(current["y1"], line["y1"])
                current["size"] = min(current["size"], line["size"])
                current["bold"] = bool(current["bold"] and line["bold"])
                current["last_line"] = line
                current["line_count"] += 1

        if current:
            paragraphs.append(current)

        blocks = []
        for paragraph in paragraphs:
            line_widths = paragraph.get("line_widths", []) or [float(paragraph["x1"] - paragraph["x0"])]
            line_heights = paragraph.get("line_heights", []) or [float(paragraph["size"])]
            avg_line_height = sum(line_heights) / max(len(line_heights), 1)
            max_line_width = max(line_widths)

            blocks.append(
                {
                    "text": paragraph["text"].strip(),
                    "visual_text": "\n".join(paragraph.get("lines", [])).strip(),
                    "lines": list(paragraph.get("lines", [])),
                    "x0": float(paragraph["x0"]),
                    "y0": float(paragraph["y0"]),
                    "x1": float(paragraph["x1"]),
                    "y1": float(paragraph["y1"]),
                    "size": float(paragraph["size"]),
                    "bold": bool(paragraph["bold"]),
                    "kind": "paragraph",
                    "line_count": int(paragraph["line_count"]),
                    "avg_line_height": float(avg_line_height),
                    "max_line_width": float(max_line_width),
                }
            )

        paragraph_pages.append((pw, ph, blocks))

    return paragraph_pages

def _extract_text_blocks(pdf_path: str, mode: str = "paragraphs"):
    """
    Extrae bloques de texto con posición usando pdfplumber.

    Parámetros:
        mode="paragraphs"
            Agrupa líneas en párrafos editables. Es el nuevo modo recomendado.

        mode="lines"
            Mantiene el comportamiento original: una caja de texto por línea.
            Útil cuando se necesita máxima fidelidad posicional.
    """
    line_pages = _extract_text_lines(pdf_path)

    if mode == "lines":
        return line_pages

    return _group_lines_into_paragraphs(line_pages)


def _extract_pdf_image_regions(pdf_path: str, dpi: int = 150):
    """
    Extrae regiones de imagen visibles del PDF como objetos independientes.

    Se usa PyMuPDF porque permite localizar imágenes con coordenadas dentro de
    cada página. Para mantener la apariencia real, cada región se rasteriza como
    PNG usando un recorte exacto del PDF, no como página completa.

    Retorno:
        [
            {
                "page_index": 0,
                "x0": ..., "y0": ..., "x1": ..., "y1": ...,
                "png": bytes,
            },
            ...
        ]
    """
    import fitz  # PyMuPDF

    regions_by_page = []
    matrix = fitz.Matrix(dpi / 72.0, dpi / 72.0)

    with fitz.open(pdf_path) as doc:
        for page_index, page in enumerate(doc):
            page_regions = []
            page_area = float(page.rect.width * page.rect.height)
            seen = set()

            try:
                image_infos = page.get_image_info(xrefs=True)
            except TypeError:
                image_infos = page.get_image_info()

            for info in image_infos:
                bbox = info.get("bbox")
                if not bbox:
                    continue

                rect = fitz.Rect(bbox)
                rect = rect & page.rect

                # Filtrar imágenes diminutas, máscaras o residuos invisibles.
                if rect.width < 8 or rect.height < 8:
                    continue
                if (rect.width * rect.height) < 96:
                    continue

                # Deduplicación por coordenadas redondeadas.
                key = (
                    round(rect.x0, 1), round(rect.y0, 1),
                    round(rect.x1, 1), round(rect.y1, 1),
                    info.get("xref"),
                )
                if key in seen:
                    continue
                seen.add(key)

                # Si la imagen cubre casi toda la página es el fondo de página
                # completo (página escaneada o fondo plano). Se omite para no
                # duplicar toda la página como objeto independiente.
                if page_area and (rect.width * rect.height) / page_area > 0.995:
                    continue

                try:
                    pix = page.get_pixmap(matrix=matrix, clip=rect, alpha=False)
                    png_bytes = pix.tobytes("png")
                except Exception as e:
                    log.debug("Página %d: imagen omitida (%s)", page_index + 1, e)
                    continue

                page_regions.append(
                    {
                        "page_index": page_index,
                        "x0": float(rect.x0),
                        "y0": float(rect.y0),
                        "x1": float(rect.x1),
                        "y1": float(rect.y1),
                        "png": png_bytes,
                    }
                )

            regions_by_page.append(page_regions)

    return regions_by_page


def _add_text_blocks_to_slide(slide, blocks, pw, ph, lft, top, iw, ih, slide_w,
                              text_mode: str, composition_mode: str,
                              preserve_line_breaks: bool):
    """
    Inserta bloques de texto en una diapositiva.

    En modo fondo fiel, el texto se mantiene blanco sobre la imagen base, como
    en el MVP original. En modo objetos independientes, no hay imagen base de
    página, por lo que el texto se inserta visible en negro.
    """
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor

    if not blocks or pw == 0 or ph == 0:
        return

    sx = iw / (pw * 12700)
    sy = ih / (ph * 12700)
    for blk in blocks:
        bx = lft + int(float(blk["x0"]) * 12700 * sx)
        by = top + int(float(blk["y0"]) * 12700 * sy)
        block_kind = str(blk.get("kind", text_mode))
        is_paragraph = block_kind == "paragraph"
        line_count = max(int(blk.get("line_count", 1)), 1)
        source_width_pt = max(float(blk.get("max_line_width", 0)), float(blk["x1"] - blk["x0"]))
        source_height_pt = max(float(blk["y1"] - blk["y0"]), float(blk.get("avg_line_height", blk["size"])) * line_count)

        if is_paragraph:
            # v3.3: ancho correlacionado directamente con el párrafo original.
            # Se evita agrandarlo de forma excesiva para que PowerPoint no cambie
            # demasiado la estructura visual del texto.
            bw = max(int(source_width_pt * 12700 * sx * 1.045), 120000)
            bh = max(int(source_height_pt * 12700 * sy * 1.22), 70000)
        else:
            bw = max(int((float(blk["x1"]) - float(blk["x0"])) * 12700 * sx * 1.08), 60000)
            bh = max(int((float(blk["y1"]) - float(blk["y0"])) * 12700 * sy * 1.10), 36000)

        # No dejar que la caja salga de la diapositiva por el lado derecho.
        max_available_w = max(slide_w - bx - 25000, 60000)
        bw = min(bw, max_available_w)

        tb = slide.shapes.add_textbox(Emu(bx), Emu(by), Emu(bw), Emu(bh))
        tf = tb.text_frame
        tf.word_wrap = True if is_paragraph else False

        # Márgenes internos mínimos para que la posición PDF → PPTX sea precisa.
        tf.margin_left = Emu(0)
        tf.margin_right = Emu(0)
        tf.margin_top = Emu(0)
        tf.margin_bottom = Emu(0)

        p = tf.paragraphs[0]
        run = p.add_run()

        if is_paragraph and preserve_line_breaks:
            # Una sola caja por párrafo, pero con saltos internos originales.
            # Esto mantiene el layout casi idéntico y evita cientos de mini cajas.
            run.text = str(blk.get("visual_text") or blk.get("text", ""))
        else:
            run.text = str(blk.get("text", ""))

        run.font.size = Pt(float(blk.get("size", 11) or 11))
        run.font.bold = bool(blk.get("bold", False))

        if composition_mode == "objects":
            # En modo objetos no hay fondo de página: texto visible.
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)  # type: ignore
        else:
            # En modo fondo fiel, se conserva la estrategia anterior.
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)  # type: ignore

        tb.fill.background()


def fase2_convert(src_path: str, dst_path: str, dpi: int = 150,
                  include_text: bool = True, text_mode: str = "paragraphs",
                  composition_mode: str = "background",
                  preserve_line_breaks: bool = True,
                  progress_cb=None) -> bool:
    from pypdf import PdfReader
    from pdf2image import convert_from_path
    from pptx import Presentation
    from pptx.util import Emu
    import io

    src = Path(src_path)
    reader = PdfReader(str(src))
    n = len(reader.pages)
    log.debug("Fase 2: %s — %d páginas, composición=%s, texto=%s (%s)",
              src.name, n, composition_mode, include_text, text_mode)

    if progress_cb:
        progress_cb(5, "Leyendo PDF…")

    text_pages = []
    if include_text:
        if progress_cb:
            progress_cb(8, "Extrayendo texto y agrupando bloques…")
        try:
            text_pages = _extract_text_blocks(src_path, mode=text_mode)
        except Exception as e:
            log.warning("Fase 2: extracción de texto omitida en %s: %s",
                        src.name, e, exc_info=True)
            if progress_cb:
                progress_cb(8, f"Extracción omitida: {e}")

    images = []
    image_regions = []

    if composition_mode == "objects":
        if progress_cb:
            progress_cb(12, "Extrayendo imágenes independientes del PDF…")
        try:
            image_regions = _extract_pdf_image_regions(src_path, dpi=dpi)
        except Exception as e:
            log.warning("Fase 2: extracción de imágenes omitida en %s: %s",
                        src.name, e, exc_info=True)
            image_regions = [[] for _ in range(n)]
            if progress_cb:
                progress_cb(12, f"Extracción de imágenes omitida: {e}")
    else:
        if progress_cb:
            progress_cb(12, f"Renderizando {n} páginas a {dpi} DPI…")
        kwargs: dict = {"dpi": dpi}
        if POPPLER_PATH:
            kwargs["poppler_path"] = POPPLER_PATH
        images = convert_from_path(str(src), **kwargs)

    prs = Presentation()
    prs.slides._sldIdLst.clear()  # type: ignore
    blank = prs.slide_layouts[6]

    for i, pdf_page in enumerate(reader.pages):
        if progress_cb:
            progress_cb(12 + int(80 * ((i + 1) / max(n, 1))), f"Página {i+1}/{n}…")

        w_pt = float(pdf_page.mediabox.width)
        h_pt = float(pdf_page.mediabox.height)
        w_emu = int(w_pt * 12700)
        h_emu = int(h_pt * 12700)

        if i == 0:
            prs.slide_width = Emu(w_emu)
            prs.slide_height = Emu(h_emu)

        slide = prs.slides.add_slide(blank)
        slide_w = int(prs.slide_width)
        slide_h = int(prs.slide_height)

        # Si todas las páginas tienen el mismo tamaño, ratio=1. Si hay mezcla de
        # orientaciones o tamaños, se centra proporcionalmente.
        ratio = min(slide_w / w_emu, slide_h / h_emu)
        iw = int(w_emu * ratio)
        ih = int(h_emu * ratio)
        lft = (slide_w - iw) // 2
        top = (slide_h - ih) // 2

        if composition_mode == "objects":
            # ── Imágenes independientes ─────────────────────────────────────
            # No se inserta la imagen completa de la página. Solo se colocan las
            # imágenes internas detectadas: fotos, tablas JPG, gráficos raster, etc.
            page_regions = image_regions[i] if i < len(image_regions) else []
            sx = iw / (w_pt * 12700) if w_pt else 1
            sy = ih / (h_pt * 12700) if h_pt else 1

            for region in page_regions:
                rx = lft + int(float(region["x0"]) * 12700 * sx)
                ry = top + int(float(region["y0"]) * 12700 * sy)
                rw = int((float(region["x1"]) - float(region["x0"])) * 12700 * sx)
                rh = int((float(region["y1"]) - float(region["y0"])) * 12700 * sy)
                if rw <= 0 or rh <= 0:
                    continue
                buf = io.BytesIO(region["png"])
                slide.shapes.add_picture(buf, Emu(rx), Emu(ry), Emu(rw), Emu(rh))
        else:
            # ── Imagen de fondo completa ────────────────────────────────────
            # Modo de máxima fidelidad visual: el PPTX conserva el aspecto del PDF.
            if i < len(images):
                buf = io.BytesIO()
                images[i].save(buf, format="PNG")
                buf.seek(0)
                slide.shapes.add_picture(buf, Emu(lft), Emu(top), Emu(iw), Emu(ih))

        # ── Texto editable ──────────────────────────────────────────────────
        if include_text and i < len(text_pages):
            pw, ph, blocks = text_pages[i]
            _add_text_blocks_to_slide(
                slide=slide,
                blocks=blocks,
                pw=pw,
                ph=ph,
                lft=lft,
                top=top,
                iw=iw,
                ih=ih,
                slide_w=slide_w,
                text_mode=text_mode,
                composition_mode=composition_mode,
                preserve_line_breaks=preserve_line_breaks,
            )

    if progress_cb:
        progress_cb(95, "Guardando PPTX…")
    prs.save(dst_path)

    if progress_cb:
        if composition_mode == "objects":
            progress_cb(100, "✅  PPTX listo — objetos independientes")          
        else:
            progress_cb(100, "✅  PPTX listo — fondo fiel + texto editable")          
    return True


# ══════════════════════════════════════════════════════════════════════════════
#  FASE 2 (variante)  —  PDF  →  WORD  (texto corrido editable)
# ══════════════════════════════════════════════════════════════════════════════

def _clean_pdf_text_artifacts(text: str) -> str:
    """
    Limpia artefactos típicos de extracción de PDF.

    Los glifos de viñeta de fuentes simbólicas (Symbol, ZapfDingbats) no tienen
    mapeo Unicode y pdfplumber los devuelve como "(cid:127)" o similar. Se
    normalizan a una viñeta real para que el texto sea legible y divisible.
    """
    import re

    # Códigos conocidos de viñeta en fuentes simbólicas.
    text = re.sub(r"\(cid:(?:127|183|149|9679|8226)\)", "\u2022", text)
    # Cualquier otro cid sin mapeo se descarta: es un glifo no representable.
    text = re.sub(r"\(cid:\d+\)", "", text)
    return " ".join(text.split())


def _split_inline_bullets(text: str):
    """
    Divide un bloque que contiene varias viñetas en una entrada por viñeta.

    Al extraer de un PDF, varios ítems de lista pueden llegar unidos en un solo
    bloque. Word necesita un párrafo por ítem para numerarlos correctamente.
    """
    marker = "\u2022"
    if text.count(marker) < 2:
        return [text]

    parts = [p.strip() for p in text.split(marker)]
    return [p for p in parts if p]


def _looks_like_numbered_heading(text: str) -> bool:
    """
    Distingue un título numerado de un ítem de lista numerada.

    "1. Introduccion" y "2.1 Hallazgos" son títulos de sección: texto corto,
    sin puntuación final y con una sola frase. Un ítem de lista numerada suele
    ser una oración más larga.
    """
    import re

    stripped = text.strip()
    if not re.match(r"^\d+(\.\d+)*[\.)]?\s+\S", stripped):
        return False

    body = re.sub(r"^\d+(\.\d+)*[\.)]?\s+", "", stripped)
    if len(body) > 70:
        return False
    if body.endswith((".", ";", ",")):
        return False
    return True


def _classify_paragraph_style(block, body_size: float):
    """
    Decide el estilo Word de un bloque de párrafo extraído del PDF.

    No se busca replicar el PDF exactamente. La meta es producir un documento
    Word natural: encabezados como encabezados, listas como listas y cuerpo
    como texto corrido que Word pueda reacomodar libremente al editar.

    Retorna: "h1" | "h2" | "h3" | "list" | "body"
    """
    text = (block.get("text") or "").strip()
    size = float(block.get("size", body_size) or body_size)
    bold = bool(block.get("bold"))

    if not text:
        return "body"

    ratio = size / body_size if body_size else 1.0
    short = len(text) < 120
    single_line = int(block.get("line_count", 1)) <= 2

    # Un título numerado ("2. Hallazgos") no es un ítem de lista.
    if _looks_like_numbered_heading(text):
        if ratio >= 1.55:
            return "h1"
        if ratio >= 1.22:
            return "h2"
        return "h3"

    if _is_bullet_or_numbered(text):
        return "list"

    # Un bloque corto, más grande que el cuerpo, es casi siempre un título.

    if short and single_line:
        if ratio >= 1.45:
            return "h1"
        if ratio >= 1.22:
            return "h2"
        if ratio >= 1.08 or bold:
            return "h3"

    return "body"


def _strip_list_marker(text: str) -> str:
    """Quita el bullet o numeral inicial; Word aporta su propia viñeta."""
    import re
    return re.sub(r"^(\u2022|\-|\*|\d+[\.)]|[a-zA-Z][\.)])\s+", "", text.strip())


def _estimate_body_size(text_pages) -> float:
    """
    Estima el tamaño de fuente del cuerpo del documento.

    Se usa la moda aproximada ponderada por cantidad de texto, porque el cuerpo
    es lo que más caracteres aporta. Así los títulos se detectan por contraste
    contra el cuerpo real y no contra un valor fijo.
    """
    from collections import Counter

    counter: "Counter" = Counter()
    for _pw, _ph, blocks in text_pages:
        for blk in blocks:
            size = round(float(blk.get("size", 11) or 11), 1)
            counter[size] += max(len(blk.get("text") or ""), 1)

    if not counter:
        return 11.0
    return float(counter.most_common(1)[0][0])


def fase2_word_convert(src_path: str, dst_path: str,
                       include_images: bool = True,
                       image_dpi: int = 150,
                       progress_cb=None) -> bool:
    """
    Convierte un PDF a un documento Word (.docx) de texto corrido editable.

    Enfoque:
        El texto se escribe como párrafos reales de Word, no como cajas
        flotantes posicionadas. Word reacomoda el contenido al editar, igual
        que en un documento nativo. Se pierde la posición exacta del PDF y se
        gana edición natural, que es el objetivo de esta salida.

    Diferencia con fase2_convert (PPTX):
        - PPTX conserva el layout visual para corregir presentaciones.
        - DOCX conserva el contenido para reescribir y corregir texto.

    Parámetros:
        include_images  Inserta las imágenes del PDF en su orden de aparición.
        image_dpi       Resolución de rasterizado de esas imágenes.
    """
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import io

    src = Path(src_path)

    if progress_cb:
        progress_cb(5, "Leyendo PDF…")

    # Texto agrupado en párrafos: es la base del documento Word.
    text_pages = _extract_text_blocks(str(src), mode="paragraphs")

    if not text_pages:
        raise RuntimeError("No se pudo leer texto del PDF")

    body_size = _estimate_body_size(text_pages)
    log.debug("Fase 2 Word: %s — %d páginas, cuerpo ≈ %g pt",
              src.name, len(text_pages), body_size)

    if progress_cb:
        progress_cb(18, f"Texto extraído (cuerpo ≈ {body_size:g} pt)")

    # Imágenes por página, en orden de aparición vertical.
    image_pages = []
    if include_images:
        if progress_cb:
            progress_cb(25, "Extrayendo imágenes…")
        try:
            image_pages = _extract_pdf_image_regions(str(src), dpi=image_dpi)
        except Exception as e:
            log.warning("Fase 2 Word: extracción de imágenes omitida en %s: %s",
                        src.name, e, exc_info=True)
            image_pages = []

    doc = Document()

    # Estilo base del documento: cuerpo legible y ancho de página carta.
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    total_pages = len(text_pages)
    empty_doc   = True

    for page_index, (pw, ph, blocks) in enumerate(text_pages):
        if progress_cb:
            pct = 30 + int(60 * (page_index / max(total_pages, 1)))
            progress_cb(pct, f"Componiendo página {page_index + 1}/{total_pages}…")

        page_images = image_pages[page_index] if page_index < len(image_pages) else []

        # Se intercala texto e imágenes según su posición vertical en la página,
        # de modo que una foto entre dos párrafos quede entre esos párrafos.
        items = []
        for blk in blocks:
            if (blk.get("text") or "").strip():
                items.append((float(blk.get("y0", 0)), "text", blk))
        for img in page_images:
            items.append((float(img.get("y0", 0)), "image", img))

        items.sort(key=lambda item: item[0])

        for _y, kind, payload in items:
            if kind == "image":
                try:
                    width_pt = float(payload["x1"]) - float(payload["x0"])
                    # Ancho útil de página carta con márgenes por defecto.
                    width_in = min(max(width_pt / 72.0, 0.6), 6.5)
                    stream = io.BytesIO(payload["png"])
                    para = doc.add_paragraph()
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.add_run().add_picture(stream, width=Inches(width_in))
                    empty_doc = False
                except Exception as e:
                    # Una imagen ilegible no debe abortar el documento completo.
                    log.debug("Página %d: imagen no insertada en Word (%s)",
                              page_index + 1, e)
                    continue
                continue

            blk   = payload
            text  = _clean_pdf_text_artifacts(blk.get("text") or "")
            if not text:
                continue
            style = _classify_paragraph_style(blk, body_size)

            # Varias viñetas unidas en un bloque: una entrada de lista por ítem.
            pieces = _split_inline_bullets(text)
            if len(pieces) > 1:
                for piece in pieces:
                    doc.add_paragraph(_strip_list_marker(piece),
                                       style="List Bullet")
                empty_doc = False
                continue

            if style == "h1":
                para = doc.add_heading(text, level=1)
            elif style == "h2":
                para = doc.add_heading(text, level=2)
            elif style == "h3":
                para = doc.add_heading(text, level=3)
            elif style == "list":
                para = doc.add_paragraph(_strip_list_marker(text),
                                          style="List Bullet")
            else:
                para = doc.add_paragraph(text)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if blk.get("bold"):
                    for run in para.runs:
                        run.bold = True

            empty_doc = False

        # Salto de página entre páginas del PDF, salvo en la última.
        if page_index < total_pages - 1:
            doc.add_page_break()

    if empty_doc:
        doc.add_paragraph(
            "No se encontró texto embebido en este PDF. "
            "Si el documento es un escaneo de imagen, se requiere OCR."
        )

    if progress_cb:
        progress_cb(95, "Guardando DOCX…")

    doc.save(dst_path)

    if progress_cb:
        progress_cb(100, "✅  Word listo — texto corrido editable")

    return True


# ══════════════════════════════════════════════════════════════════════════════
#  UI — WIDGETS COMPARTIDOS
# ══════════════════════════════════════════════════════════════════════════════

def _open_folder(path: Path):
    folder = str(path.parent if path.is_file() else path)
    if   sys.platform == "win32":  os.startfile(folder)
    elif sys.platform == "darwin": subprocess.run(["open",      folder])
    else:                          subprocess.run(["xdg-open",  folder])
# ══════════════════════════════════════════════════════════════════════════════
#  MOTOR — RECUPERACIÓN DE TABLAS EXCEL
#  (basado en PDF Table Recovery Mini v0.1 — GoTopo 2026)
# ══════════════════════════════════════════════════════════════════════════════

import csv
import re as _re
from dataclasses import dataclass
from statistics import median as _median
from typing import List, Optional, Sequence, Tuple


@dataclass
class _Cell:
    text: str
    x0: float
    x1: float
    top: float
    bottom: float


@dataclass
class _RecoveredPage:
    page_number: int
    width: float
    height: float
    word_count: int
    rows: List[List[str]]
    method: str = "bandas"          # "rejilla" o "bandas"
    suspicious: set = None          # coordenadas (fila, col) dudosas
    audit: dict = None              # conteo por tipo de sospecha


def _parse_crop(value: Optional[str]):
    if not value:
        return None
    parts = [float(x.strip()) for x in value.split(",")]
    if len(parts) != 4:
        raise ValueError("Recorte debe tener formato x0,y0,x1,y1")
    return tuple(parts)


def _looks_numeric_financial(text: str) -> bool:
    s = text.strip()
    if not any(ch.isdigit() for ch in s):
        return False
    if "%" in s:
        return False
    return bool(_re.fullmatch(r"[₡$€()\-\s.,0-9]+", s))


def _format_thousands_spaces(digits: str) -> str:
    if not digits:
        return ""
    groups: List[str] = []
    while digits:
        groups.append(digits[-3:])
        digits = digits[:-3]
    return " ".join(reversed(groups))


def _normalize_financial_cell(text: str) -> str:
    s = " ".join(text.strip().split())
    if not _looks_numeric_financial(s):
        return s
    has_colon  = "₡" in s
    has_dollar = "$" in s
    negative   = "(" in s and ")" in s
    leading_minus = s.lstrip().startswith("-")
    digits = "".join(ch for ch in s if ch.isdigit())
    if not digits:
        return s
    number = _format_thousands_spaces(digits)
    prefix = "₡" if has_colon else "$" if has_dollar else ""
    if negative:
        return f"({prefix}{number})"
    if leading_minus:
        return f"-{prefix}{number}"
    return f"{prefix}{number}"


def _get_words_table(page, max_font_size: float, crop):
    """
    Extrae palabras de una página para reconstrucción de tabla.

    Dos decisiones importantes, ambas producto de casos reales:

    1) use_text_flow=True
       Cuando una tabla de Excel se pega en Word como metarchivo vectorial,
       el PDF posiciona cada glifo por separado. Sin este parámetro pdfplumber
       devuelve un token por carácter y el texto sale como "R e s u l t a d o".

    2) NO se pasa "size" en extra_attrs
       pdfplumber corta palabra cada vez que un atributo cambia, y en esos
       metarchivos el tamaño difiere en el octavo decimal entre glifos
       (7.145899333 vs 7.145899415). Eso bastaba para fragmentar todo.
       El filtro de tamaño se hace con la altura del token, que es estable.
    """
    target = page.crop(crop) if crop else page
    words  = target.extract_words(
        x_tolerance=2, y_tolerance=3,
        keep_blank_chars=False, use_text_flow=True,
    ) or []

    filtered = []
    for w in words:
        text = str(w.get("text", "")).strip()
        if not text:
            continue
        # Altura del token como aproximación del tamaño de fuente.
        height = float(w.get("bottom", 0)) - float(w.get("top", 0))
        if height > max_font_size * 1.35:
            continue
        filtered.append(w)
    return filtered


def _is_thousands_continuation(text: str) -> bool:
    """
    Reconoce el resto de un monto partido por ancho de celda.

    Solo acepta un grupo de millar suelto, con paréntesis de cierre opcional:
        "279", "651)", "08"
    Se exige que sean dígitos puros para no absorber etiquetas ni otro monto.
    """
    import re
    return bool(re.fullmatch(r"\d{1,3}\)?", text.strip()))


def _is_money_like(text: str) -> bool:
    """Monto o porcentaje que podría venir cortado por wrap de celda."""
    import re
    return bool(re.fullmatch(r"[\(\-]?\$?[\d.,]+\)?%?", text.strip()))


def _merge_side_thousands(words, max_gap: float = 8.0):
    """
    Une los grupos de millar que van uno al lado del otro en el mismo renglón.

    El separador de miles en estos estados financieros es un espacio, así que
    "$670 279" llega como dos tokens. Unirlos antes de agrupar columnas evita
    que cada mitad genere su propia columna y que un monto quede partido en dos
    celdas, que es justo el error que hace dudar de toda la tabla.
    """
    ordered = sorted(words, key=lambda w: (float(w["top"]), float(w["x0"])))
    consumed = set()
    result   = []

    for i, head in enumerate(ordered):
        if i in consumed:
            continue

        text   = str(head.get("text", ""))
        x1     = float(head["x1"])
        mid    = (float(head["top"]) + float(head["bottom"])) / 2
        bottom = float(head["bottom"])

        if _is_money_like(text):
            # Se admiten varios grupos seguidos: "$1 234 567".
            advanced = True
            while advanced:
                advanced = False
                for j, nxt in enumerate(ordered):
                    if j in consumed or j == i:
                        continue
                    nxt_mid = (float(nxt["top"]) + float(nxt["bottom"])) / 2
                    if abs(nxt_mid - mid) > 2.5:
                        continue
                    if not (0 <= (float(nxt["x0"]) - x1) <= max_gap):
                        continue
                    if not _is_thousands_continuation(str(nxt.get("text", ""))):
                        continue
                    text = text + " " + str(nxt["text"])
                    x1   = float(nxt["x1"])
                    bottom = max(bottom, float(nxt["bottom"]))
                    consumed.add(j)
                    advanced = True
                    break

        merged = dict(head)
        merged["text"]   = text
        merged["x1"]     = x1
        merged["bottom"] = bottom
        result.append(merged)

    return result


def _merge_wrapped_numbers(words, tol_x: float = 1.8, max_gap: float = 6.0):
    """
    Une montos partidos en dos líneas dentro de una misma celda.

    Cuando la celda de Word es más angosta que el número, "$670 279" se
    imprime en dos líneas. El agrupador por bandas las ve como dos filas
    distintas y el número queda cortado y descolocado: ese es el fallo que
    hace desconfiar de toda la conversión.

    Regla deliberadamente estricta, para no inventar datos:
        - el fragmento de abajo debe ser dígitos puros (grupo de millar)
        - debe compartir el borde derecho con el monto de arriba
          (los montos van alineados a la derecha, así que x1 coincide)
        - el hueco vertical debe ser de una línea, no de una fila
        - una sola unión por monto, sin encadenar

    Sin encadenar es esencial: de lo contrario "$670 279" seguiría absorbiendo
    el monto de la fila siguiente por compartir el mismo borde derecho.
    """
    ordered = sorted(words, key=lambda w: (float(w["top"]), float(w["x0"])))
    consumed = set()
    result   = []

    for i, upper in enumerate(ordered):
        if i in consumed:
            continue

        text = str(upper.get("text", ""))

        if _is_money_like(text):
            # Si el monto ya continúa a la derecha en el mismo renglón, está
            # completo y no hay wrap. Verificarlo evita el error inverso:
            # absorber el número de la fila de abajo cuando no hacía falta.
            upper_mid = (float(upper["top"]) + float(upper["bottom"])) / 2
            has_side_continuation = any(
                abs(((float(o["top"]) + float(o["bottom"])) / 2) - upper_mid) < 2.5
                and 0 <= (float(o["x0"]) - float(upper["x1"])) <= 8
                and _is_thousands_continuation(str(o.get("text", "")))
                for k, o in enumerate(ordered) if k != i
            )

            if not has_side_continuation:
                for j in range(i + 1, len(ordered)):
                    if j in consumed:
                        continue
                    lower = ordered[j]
                    gap   = float(lower["top"]) - float(upper["bottom"])

                    if gap > max_gap:
                        break          # ya estamos en otra fila
                    if gap < -2:
                        continue       # mismo renglón, no es continuación

                    same_right = abs(float(lower["x1"]) - float(upper["x1"])) <= tol_x
                    if not same_right:
                        continue
                    if not _is_thousands_continuation(str(lower.get("text", ""))):
                        continue

                    # El fragmento de abajo no puede ser el grupo de millar de
                    # un monto de su propia fila. Si tiene un monto pegado a su
                    # izquierda, ya pertenece a ese monto y no es continuación.
                    lower_mid = (float(lower["top"]) + float(lower["bottom"])) / 2
                    belongs_to_own_row = any(
                        abs(((float(o["top"]) + float(o["bottom"])) / 2) - lower_mid) < 2.5
                        and 0 <= (float(lower["x0"]) - float(o["x1"])) <= 8
                        and _is_money_like(str(o.get("text", "")))
                        for k, o in enumerate(ordered) if k != j
                    )
                    if belongs_to_own_row:
                        continue

                    text = text + " " + str(lower["text"])
                    consumed.add(j)
                    break          # una sola unión

        merged = dict(upper)
        merged["text"] = text
        result.append(merged)

    return result


def _group_words_into_lines_table(words, y_tol: float):
    if not words:
        return []
    sorted_words = sorted(words, key=lambda w: (float(w.get("top", 0)), float(w.get("x0", 0))))
    lines: List[List[dict]] = []
    for w in sorted_words:
        top    = float(w.get("top", 0))
        placed = False
        for line in lines:
            line_top = _median(float(x.get("top", 0)) for x in line)
            if abs(top - line_top) <= y_tol:
                line.append(w)
                placed = True
                break
        if not placed:
            lines.append([w])
    for line in lines:
        line.sort(key=lambda w: float(w.get("x0", 0)))
    lines.sort(key=lambda line: _median(float(w.get("top", 0)) for w in line))
    return lines


def _words_to_cells_table(line, gap_tol: float) -> List[_Cell]:
    if not line:
        return []
    cells: List[_Cell] = []
    cur_words: List[dict] = []

    def flush():
        if not cur_words:
            return
        text = " ".join(str(w["text"]) for w in cur_words)
        text = _normalize_financial_cell(text)
        cells.append(_Cell(
            text   = text,
            x0     = float(min(w["x0"]     for w in cur_words)),
            x1     = float(max(w["x1"]     for w in cur_words)),
            top    = float(min(w.get("top",    0) for w in cur_words)),
            bottom = float(max(w.get("bottom", 0) for w in cur_words)),
        ))
        cur_words.clear()

    prev_x1 = None
    for w in sorted(line, key=lambda x: float(x.get("x0", 0))):
        x0 = float(w["x0"])
        if prev_x1 is not None and (x0 - prev_x1) > gap_tol:
            flush()
        cur_words.append(w)
        prev_x1 = float(w["x1"])
    flush()
    return cells


# ══════════════════════════════════════════════════════════════════════════════
#  FASE 3  —  RECONSTRUCCIÓN POR REJILLA
# ══════════════════════════════════════════════════════════════════════════════

def _collapse_edges(values, tol: float = 3.0):
    """Agrupa coordenadas casi iguales en un solo borde."""
    edges = []
    for v in sorted(values):
        if not edges or (v - edges[-1]) > tol:
            edges.append(v)
    return edges


def _collect_grid_edges(page, min_span: float = 25.0):
    """
    Deriva los bordes de columna y fila a partir de la rejilla dibujada en el PDF.

    Dos orígenes posibles, según cómo se creó la tabla:
        - Rectángulos: bordes de celda de una tabla de Word.
        - Líneas: cuadrícula de Excel cuando se pegó como metarchivo vectorial.

    Retorna (xs, ys) o (None, None) si el PDF no trae rejilla utilizable.
    """
    xs, ys = set(), set()

    for rect in getattr(page, "rects", []) or []:
        try:
            x0, x1 = float(rect["x0"]), float(rect["x1"])
            y0, y1 = float(rect["top"]), float(rect["bottom"])
        except (KeyError, TypeError, ValueError):
            continue
        if (x1 - x0) >= 2 or (y1 - y0) >= 2:
            xs.update((round(x0, 1), round(x1, 1)))
            ys.update((round(y0, 1), round(y1, 1)))

    # Las cuadrículas de Excel suelen venir dibujadas en segmentos cortos, uno
    # por celda. Exigir que cada segmento sea largo descarta casi toda la
    # rejilla, así que se acumula la "tinta" por coordenada y se conserva la
    # coordenada cuando el total suma una extensión significativa.
    ink_x: dict = {}
    ink_y: dict = {}

    for source in ("lines", "curves"):
        for item in getattr(page, source, []) or []:
            try:
                x0, x1 = float(item["x0"]), float(item["x1"])
                y0, y1 = float(item["top"]), float(item["bottom"])
            except (KeyError, TypeError, ValueError):
                continue

            if abs(x1 - x0) < 1.5 and abs(y1 - y0) >= 1.5:            # vertical
                key = round((x0 + x1) / 2, 0)
                ink_x[key] = ink_x.get(key, 0.0) + abs(y1 - y0)
            elif abs(y1 - y0) < 1.5 and abs(x1 - x0) >= 1.5:          # horizontal
                key = round((y0 + y1) / 2, 0)
                ink_y[key] = ink_y.get(key, 0.0) + abs(x1 - x0)

    # Umbral de tinta: descarta marcas aisladas y conserva bordes reales.
    xs.update(k for k, total in ink_x.items() if total >= 12)
    ys.update(k for k, total in ink_y.items() if total >= 12)

    xs = _collapse_edges(xs)
    ys = _collapse_edges(ys)

    # Una rejilla utilizable necesita al menos 3 columnas y 3 filas,
    # y cubrir una extensión razonable de la página.
    if len(xs) < 4 or len(ys) < 4:
        return None, None
    if (xs[-1] - xs[0]) < min_span or (ys[-1] - ys[0]) < min_span:
        return None, None

    return xs, ys


def _band_index(edges, value: float) -> int:
    """Índice de la banda de la rejilla en la que cae un valor."""
    index = -1
    for i, edge in enumerate(edges):
        if value >= edge:
            index = i
        else:
            break
    if index >= len(edges) - 1:
        return len(edges) - 2
    return index


def _cluster_column_bounds(words, tol: float = 6.0):
    """
    Deriva los límites de columna a partir de los montos, no de todo el texto.

    En una tabla financiera las cifras van alineadas a la derecha y forman
    columnas muy definidas, mientras que las etiquetas son de ancho irregular
    y romperían el agrupamiento en decenas de columnas falsas.

    Por eso se agrupan los bordes derechos solo de las celdas numéricas, y todo
    lo que quede a la izquierda de la primera cifra se trata como una única
    columna de etiquetas.
    """
    # Solo montos con cuerpo real definen columna. Tokens como "26" de un
    # encabezado "enero-26" partido en dos líneas crearían columnas falsas.
    def _is_real_amount(text: str) -> bool:
        s = text.strip()
        if not _looks_numeric_financial(s):
            return False
        digits = sum(1 for ch in s if ch.isdigit())
        return ("$" in s) or ("," in s) or digits >= 4

    numeric = [w for w in words if _is_real_amount(str(w.get("text", "")))]
    source  = numeric if len(numeric) >= 6 else words

    edges = []
    for x in sorted(float(w["x1"]) for w in source):
        if not edges or (x - edges[-1][-1]) > tol:
            edges.append([x])
        else:
            edges[-1].append(x)

    centers = [sum(group) / len(group) for group in edges]
    if len(centers) < 2:
        return None

    # Divisiones a mitad de camino entre columnas contiguas.
    splits = [(centers[i] + centers[i + 1]) / 2 for i in range(len(centers) - 1)]

    # Frontera de la columna de etiquetas: justo antes de la primera cifra.
    if numeric and len(numeric) >= 6:
        first_number_x0 = min(float(w["x0"]) for w in numeric)
        label_edge = first_number_x0 - 2
        if all(label_edge < s for s in splits):
            splits.insert(0, label_edge)

    return [float("-inf")] + splits + [float("inf")]


def _recover_page_grid(page, page_number: int, min_words: int,
                        max_font_size: float, crop):
    """
    Reconstruye la tabla combinando dos fuentes de información.

    Columnas: los bordes verticales dibujados en el PDF, cuando existen y son
    coherentes. Es lo más fiel, porque no estima dónde empieza cada columna.
    Si el PDF no trae rejilla usable, se agrupan los bordes derechos del texto.

    Filas: las líneas de texto. La rejilla horizontal resultó poco confiable en
    la práctica, porque muchas hojas de Excel solo dibujan separadores de grupo
    y no una línea por fila.

    Antes de todo se unen los montos partidos por ancho de celda, que es el
    fallo que descoloca los datos y hace desconfiar de la conversión completa.

    Retorna (rows, n_columnas) o (None, 0) si la página no da para tabla.
    """
    words = _get_words_table(page, max_font_size=max_font_size, crop=crop)
    if len(words) < min_words:
        return None, 0

    xs, ys = _collect_grid_edges(page)

    # Recortar al área de la tabla: deja fuera títulos y párrafos del documento.
    if xs and ys:
        # Solo se recorta en vertical. La columna de etiquetas suele quedar a la
        # izquierda del primer borde de la rejilla, así que recortar en X la
        # eliminaría junto con el texto del documento.
        y_min, y_max = ys[0] - 3, ys[-1] + 3
        inside = [
            w for w in words
            if y_min <= ((float(w["top"]) + float(w["bottom"])) / 2) <= y_max
            and float(w["x1"]) <= xs[-1] + 6
        ]
        if len(inside) >= min_words * 0.4:
            words = inside

    # Primero los grupos de millar contiguos, luego los montos partidos por
    # wrap. El orden importa: unir lateralmente primero deja montos completos
    # y evita confundir un grupo de millar con una continuación de otra fila.
    words = _merge_side_thousands(words)
    words = _merge_wrapped_numbers(words)

    # ── Límites de columna ───────────────────────────────────────────────────
    # Una rejilla con demasiadas divisiones suele venir de bordes internos de
    # Word y parte las etiquetas; en ese caso se prefiere el agrupamiento.
    bounds = None
    if xs and 3 <= (len(xs) - 1) <= 15:
        # Todos los bordes se usan como divisiones, y se agregan extremos
        # abiertos: el de la izquierda recoge la columna de etiquetas.
        bounds = [float("-inf")] + list(xs) + [float("inf")]

    if bounds is None:
        bounds = _cluster_column_bounds(words)
    if not bounds:
        return None, 0

    n_cols = len(bounds) - 1

    # ── Filas desde las líneas de texto ──────────────────────────────────────
    lines = _group_words_into_lines_table(words, y_tol=4.0)
    if not lines:
        return None, 0

    rows_out: List[List[str]] = []

    for line in lines:
        cells = [[] for _ in range(n_cols)]
        for w in line:
            cx = (float(w["x0"]) + float(w["x1"])) / 2
            idx = _band_index(bounds, cx)
            idx = max(0, min(idx, n_cols - 1))
            cells[idx].append((float(w["x0"]), str(w["text"])))

        row = []
        for fragments in cells:
            if not fragments:
                row.append("")
                continue
            text = " ".join(t for _x, t in sorted(fragments))
            row.append(_normalize_financial_cell(text))

        while row and not row[-1].strip():
            row.pop()

        if any(c.strip() for c in row):
            rows_out.append(row)

    # ── Compactar filas partidas por wrap desigual ───────────────────────────
    # Cuando unas celdas de la fila hacen wrap y otras no, las que caben quedan
    # centradas y aterrizan en una línea de texto aparte. El valor es correcto
    # pero aparece en una fila extra, lo que da la impresión de dato descolocado.
    # Solo se fusiona cuando no hay ninguna columna ocupada en ambas filas, así
    # que nunca se sobrescribe un dato con otro.
    compacted: List[List[str]] = []
    for row in rows_out:
        filled = [i for i, c in enumerate(row) if c.strip()]

        if compacted and len(filled) <= 3:
            previous = compacted[-1]
            prev_filled = {i for i, c in enumerate(previous) if c.strip()}

            # La fila dispersa no puede traer etiqueta propia ni chocar.
            has_label = bool(row) and bool(row[0].strip())
            collision = any(i in prev_filled for i in filled)

            if filled and not collision and not has_label and len(prev_filled) >= 2:
                while len(previous) < len(row):
                    previous.append("")
                for i in filled:
                    previous[i] = row[i]
                continue

        compacted.append(list(row))

    rows_out = compacted

    # ── Descartar filas de prosa del documento ───────────────────────────────
    # Una fila de tabla tiene varias celdas; un párrafo cae todo en pocas.
    filtered = []
    for row in rows_out:
        filled  = sum(1 for c in row if c.strip())
        numeric = sum(1 for c in row if _looks_numeric_financial(c))
        if filled >= 3 or numeric >= 2:
            filtered.append(row)

    if len(filtered) < 3:
        return None, 0

    return filtered, n_cols


# ══════════════════════════════════════════════════════════════════════════════
#  FASE 3  —  AUDITORÍA DE INTEGRIDAD
# ══════════════════════════════════════════════════════════════════════════════

def _audit_recovered_rows(rows):
    """
    Marca las celdas que podrían contener un número incompleto.

    El riesgo real de esta función no es equivocarse: es equivocarse en
    silencio. Un monto partido que queda descolocado hace que el usuario
    pierda la confianza en toda la tabla, y lo sensato entonces es descartar
    la conversión completa.

    Por eso se señalan las celdas dudosas en el Excel, en lugar de suponer
    que todo salió bien. Señales tomadas como sospechosas:

        - Un grupo de millar suelto ("279") sin símbolo ni separador: casi
          siempre es la segunda mitad de un monto partido.
        - Una celda con dos o más montos pegados: columnas que se fusionaron.

    Retorna (coordenadas_sospechosas, resumen_por_tipo).
    """
    import re

    suspicious = set()
    summary = {"fragmento": 0, "fusionada": 0}

    orphan_group = re.compile(r"^\d{1,3}\)?$")
    money_token  = re.compile(r"[\(\-]?\$\s?[\d.,\s]+\)?")

    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            text = (cell or "").strip()
            if not text:
                continue

            # Grupo de millar suelto: probable mitad de un monto.
            if orphan_group.match(text):
                suspicious.add((r, c))
                summary["fragmento"] += 1
                continue

            # Varios montos en una sola celda: columnas fusionadas.
            if len(money_token.findall(text)) >= 2:
                suspicious.add((r, c))
                summary["fusionada"] += 1
                continue

            # Varios porcentajes en una sola celda.
            if text.count("%") >= 2:
                suspicious.add((r, c))
                summary["fusionada"] += 1

    return suspicious, summary


def _recover_page_table(page, page_number: int, min_words: int,
                         max_font_size: float, y_tol: float,
                         gap_tol: float, crop,
                         prefer_grid: bool = True) -> Optional[_RecoveredPage]:
    """
    Recupera la tabla de una página con dos estrategias en orden de calidad.

    1) Rejilla — usa los bordes de celda dibujados en el PDF. Es el método
       fiel: no adivina dónde empieza cada columna.
    2) Bandas — agrupa por altura y separa por huecos horizontales. Es el
       comportamiento histórico y sirve cuando el PDF no trae rejilla.

    El método usado queda registrado en el resultado, porque al usuario le
    importa saber si la tabla vino de la rejilla real o de una estimación.
    """
    words = _get_words_table(page, max_font_size=max_font_size, crop=crop)
    if len(words) < min_words:
        return None

    rows: Optional[List[List[str]]] = None
    method = "bandas"

    # ── Estrategia 1: rejilla ────────────────────────────────────────────────
    if prefer_grid:
        try:
            grid_rows, _cols = _recover_page_grid(
                page,
                page_number=page_number,
                min_words=min_words,
                max_font_size=max_font_size,
                crop=crop,
            )
            if grid_rows:
                rows = grid_rows
                method = "rejilla"
        except Exception as e:
            log.warning("Página %d: recuperación por rejilla falló (%s); usando bandas",
                        page_number, e)
            log.debug("Traceback de rejilla, página %d", page_number, exc_info=True)
            rows = None

    # ── Estrategia 2: bandas (respaldo) ──────────────────────────────────────
    if rows is None:
        merged = _merge_wrapped_numbers(_merge_side_thousands(words))
        lines  = _group_words_into_lines_table(merged, y_tol=y_tol)
        rows   = []
        for line in lines:
            cells = _words_to_cells_table(line, gap_tol=gap_tol)
            if cells:
                rows.append([c.text for c in cells])

        table_like = sum(1 for r in rows if len(r) >= 2)
        if table_like < 5:
            log.debug("Página %d: sin tabla (%d filas tabulares)", page_number, table_like)
            return None

    suspicious, audit = _audit_recovered_rows(rows)
    log.debug("Página %d: tabla por %s — %d filas, %d celdas dudosas",
              page_number, method, len(rows), len(suspicious or ()))

    return _RecoveredPage(
        page_number=page_number,
        width=float(page.width),
        height=float(page.height),
        word_count=len(words),
        rows=rows,
        method=method,
        suspicious=suspicious,
        audit=audit,
    )


def _export_csv_tables(pages: List[_RecoveredPage], out_dir: Path) -> List[Path]:
    out_paths = []
    for rp in pages:
        path = out_dir / f"pagina_{rp.page_number:03d}_tabla.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rp.rows)
        out_paths.append(path)
    return out_paths


def _export_xlsx_tables(pages: List[_RecoveredPage], out_path: Path) -> Optional[Path]:
    """
    Exporta las tablas recuperadas a Excel, señalando lo que conviene revisar.

    Las celdas dudosas se pintan en ámbar y llevan un comentario. La idea es
    que el usuario sepa exactamente dónde mirar en lugar de tener que verificar
    la tabla entera o descartarla por desconfianza.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.comments import Comment
        from openpyxl.utils import get_column_letter
    except Exception as e:
        log.error("openpyxl no disponible (%s): no se generará el Excel", e)
        return None

    wb = Workbook()
    thin         = Side(style="thin", color="DDDDDD")
    header_fill  = PatternFill("solid", fgColor="E8F1FF")
    warn_fill    = PatternFill("solid", fgColor="FFE4B5")   # ámbar: revisar
    grid_fill    = PatternFill("solid", fgColor="E8F5EE")   # verde: rejilla
    band_fill    = PatternFill("solid", fgColor="FFF4E0")   # crema: bandas

    if pages and wb.active:
        wb.remove(wb.active)

    total_warn = 0

    for rp in pages:
        ws_name = _re.sub(r"[\\/*?:\[\]]", "_", f"Pagina_{rp.page_number}")[:31]
        ws = wb.create_sheet(ws_name)

        suspicious = rp.suspicious or set()
        audit      = rp.audit or {}
        total_warn += len(suspicious)

        # ── Encabezado informativo ───────────────────────────────────────────
        metodo_txt = ("Rejilla del PDF (fiel)" if rp.method == "rejilla"
                      else "Bandas de texto (estimado)")
        ws.append([f"Página {rp.page_number}", f"Método: {metodo_txt}"])
        ws.cell(1, 1).font = Font(bold=True, size=12)
        ws.cell(1, 2).font = Font(bold=True)
        ws.cell(1, 2).fill = grid_fill if rp.method == "rejilla" else band_fill

        if suspicious:
            ws.append([
                f"⚠ {len(suspicious)} celda(s) a revisar",
                f"Posibles números partidos: {audit.get('fragmento', 0)}  |  "
                f"Posibles columnas fusionadas: {audit.get('fusionada', 0)}",
            ])
            ws.cell(2, 1).font = Font(bold=True, color="B36B00")
            ws.cell(2, 2).font = Font(color="B36B00")
        else:
            ws.append(["✓ Sin celdas dudosas detectadas", ""])
            ws.cell(2, 1).font = Font(bold=True, color="107C41")

        # Fila separadora. No se usa ws.max_row para ubicar los datos porque
        # una fila vacía avanza el cursor sin registrar celdas, y eso desplaza
        # el marcado de las celdas dudosas una fila hacia arriba.
        ws.append([""])
        data_start = 4

        # ── Datos ────────────────────────────────────────────────────────────
        for row in rp.rows:
            ws.append(row)

        max_cols = max((len(r) for r in rp.rows), default=1)

        # ── Marcado de celdas dudosas ────────────────────────────────────────
        for (r_idx, c_idx) in suspicious:
            cell = ws.cell(data_start + r_idx, c_idx + 1)
            cell.fill = warn_fill
            cell.font = Font(bold=True, color="8B4500")
            try:
                cell.comment = Comment(
                    "Revisar: este valor puede estar incompleto o fusionado "
                    "con otra columna. Compare contra el PDF original.",
                    "PDF Quick Disassembler",
                )
            except Exception:
                pass

        # ── Bordes y ancho de columna ────────────────────────────────────────
        for col_idx in range(1, max_cols + 1):
            width = 8
            for row_idx in range(data_start, ws.max_row + 1):
                cell = ws.cell(row_idx, col_idx)
                if cell.value is not None:
                    width = max(width, min(len(str(cell.value)) + 2, 42))
                cell.border = Border(top=thin, left=thin,
                                      right=thin, bottom=thin)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # ── Primeras filas de datos como encabezado visual ───────────────────
        for row_idx in range(data_start, min(ws.max_row, data_start + 1) + 1):
            for col_idx in range(1, max_cols + 1):
                cell = ws.cell(row_idx, col_idx)
                if (row_idx - data_start, col_idx - 1) in suspicious:
                    continue
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", wrap_text=True)

        ws.freeze_panes = ws.cell(data_start + 1, 1).coordinate

    # ── Hoja resumen cuando hay varias páginas ───────────────────────────────
    if len(pages) > 1:
        summary = wb.create_sheet("Resumen", 0)
        summary.append(["Página", "Método", "Filas", "Celdas a revisar"])
        for c in range(1, 5):
            summary.cell(1, c).font = Font(bold=True)
            summary.cell(1, c).fill = header_fill
        for rp in pages:
            summary.append([
                rp.page_number,
                "Rejilla del PDF" if rp.method == "rejilla" else "Bandas (estimado)",
                len(rp.rows),
                len(rp.suspicious or set()),
            ])
        for col, width in zip("ABCD", (10, 22, 10, 18)):
            summary.column_dimensions[col].width = width

    wb.save(out_path)
    return out_path


def tabla_excel_recover(src_path: str, out_dir: Path,
                         max_font_size: float = 16.0,
                         y_tol: float = 4.0,
                         gap_tol: float = 18.0,
                         min_words: int = 80,
                         crop_str: Optional[str] = None,
                         progress_cb=None) -> dict:
    """
    Orquestador de recuperación de tablas Excel desde PDF.
    Retorna dict con claves: recovered (int), xlsx (Path|None), csvs (list), msg (str)
    """
    import pdfplumber

    crop = _parse_crop(crop_str) if crop_str else None
    recovered: List[_RecoveredPage] = []

    if progress_cb: progress_cb(5, "Abriendo PDF…")

    with pdfplumber.open(src_path) as pdf:
        n = len(pdf.pages)
        log.debug("Table Recovery: %s — %d páginas", Path(src_path).name, n)
        if progress_cb: progress_cb(10, f"Analizando {n} páginas…")

        for i, page in enumerate(pdf.pages):
            pct = 10 + int(75 * (i / max(n, 1)))
            if progress_cb: progress_cb(pct, f"Escaneando página {i+1}/{n}…")

            rp = _recover_page_table(
                page=page,
                page_number=i + 1,
                min_words=min_words,
                max_font_size=max_font_size,
                y_tol=y_tol,
                gap_tol=gap_tol,
                crop=crop,
            )
            if rp:
                recovered.append(rp)

    if not recovered:
        if progress_cb: progress_cb(100, "⚠️  Sin tablas compatibles encontradas")
        return {"recovered": 0, "xlsx": None, "csvs": [], "msg": "sin_tablas",
                "grid_pages": 0, "band_pages": 0, "warn_cells": 0}

    out_dir.mkdir(parents=True, exist_ok=True)

    if progress_cb: progress_cb(88, "Exportando CSV…")
    csvs = _export_csv_tables(recovered, out_dir)

    if progress_cb: progress_cb(94, "Exportando Excel…")
    stem  = Path(src_path).stem
    xlsx  = _export_xlsx_tables(recovered, out_dir / f"{stem}_tablas_recuperadas.xlsx")

    # Resumen de calidad: al usuario le importa saber si puede confiar
    # en la tabla o si hay celdas que debe verificar contra el PDF.
    grid_pages = sum(1 for rp in recovered if rp.method == "rejilla")
    warn_cells = sum(len(rp.suspicious or set()) for rp in recovered)

    if progress_cb:
        if warn_cells:
            progress_cb(100, f"✅  {len(recovered)} tabla(s)  —  "
                              f"{warn_cells} celda(s) marcada(s) para revisar")
        else:
            progress_cb(100, f"✅  {len(recovered)} tabla(s) recuperadas sin dudas")

    return {
        "recovered":   len(recovered),
        "xlsx":        xlsx,
        "csvs":        csvs,
        "msg":         "ok",
        "grid_pages":  grid_pages,
        "band_pages":  len(recovered) - grid_pages,
        "warn_cells":  warn_cells,
    }




# ── Interfaz gráfica: solo modo desktop ──
if _TKINTER_OK:


    class DropZone(Frame):
        def __init__(self, parent, allowed_ext: list, label: str,
                     on_file=None, **kw):
            super().__init__(parent, bg=C_CARD,
                             highlightbackground=C_BORDER,
                             highlightthickness=2, **kw)
            self.allowed_ext = [e.lower() for e in allowed_ext]
            self.on_file     = on_file
            self._path       = ""
            self._build(label)
            self._try_dnd()

        def _build(self, label: str):
            inner = Frame(self, bg=C_CARD, padx=22, pady=14)
            inner.pack(fill="both", expand=True)

            Label(inner, text="📂", font=("Segoe UI Emoji", 26),
                  bg=C_CARD, fg=C_ACCENT).pack()
            Label(inner, text=label,
                  font=("Segoe UI", 11, "bold"),
                  bg=C_CARD, fg=C_TEXT).pack(pady=(4, 1))
            Label(inner,
                  text="  •  ".join(e.upper() for e in self.allowed_ext),
                  font=("Segoe UI", 9), bg=C_CARD, fg=C_MUTED).pack()
            Label(inner, text="Arrastra aquí  o",
                  font=("Segoe UI", 9), bg=C_CARD, fg=C_MUTED).pack(pady=(8, 4))

            Button(inner, text="  Seleccionar archivo  ",
                   font=("Segoe UI", 10, "bold"),
                   bg=C_ACCENT, fg="white", activebackground="#3A6AEE",
                   relief="flat", padx=16, pady=6, cursor="hand2",
                   command=self._browse).pack()

            self.lbl = Label(inner, text="",
                              font=("Segoe UI", 9, "bold"),
                              bg=C_CARD, fg=C_GREEN, wraplength=400)
            self.lbl.pack(pady=(8, 0))

        def _try_dnd(self):
            try:
                self.drop_target_register("DND_Files")   # type: ignore
                self.dnd_bind("<<Drop>>", lambda e: self._set(e.data.strip().strip("{}")))  # type: ignore
                self.config(highlightbackground=C_ACCENT)
            except Exception:
                pass

        def _browse(self):
            flt = [(f"{e.upper()} files", f"*{e}") for e in self.allowed_ext]
            flt.append(("Todos los archivos", "*.*"))
            p = filedialog.askopenfilename(filetypes=flt)
            if p: self._set(p)

        def _set(self, path: str):
            if Path(path).suffix.lower() not in self.allowed_ext:
                messagebox.showerror("Formato no válido",
                    f"Selecciona: {', '.join(e.upper() for e in self.allowed_ext)}")
                return
            self._path = path
            self.lbl.config(text=f"✅  {Path(path).name}")
            self.config(highlightbackground=C_GREEN)
            if self.on_file: self.on_file(path)

        def get(self) -> str: return self._path
        def clear(self):
            self._path = ""
            self.lbl.config(text="")
            self.config(highlightbackground=C_BORDER)


    class StatusBar(Frame):
        def __init__(self, parent, **kw):
            super().__init__(parent, bg=C_PANEL, **kw)
            self._msg = StringVar(value="Listo")
            Label(self, textvariable=self._msg, bg=C_PANEL, fg=C_MUTED,
                  font=("Segoe UI", 9)).pack(anchor="w", padx=12, pady=(4, 0))
            self.bar = Progressbar(self, mode="determinate", maximum=100)
            self.bar.pack(fill="x", padx=12, pady=(2, 6))

        def set_status(self, pct: int, msg: str):
            self.bar["value"] = pct
            self._msg.set(msg)
            self.update_idletasks()


    # ══════════════════════════════════════════════════════════════════════════════
    #  TAB FASE 1
    # ══════════════════════════════════════════════════════════════════════════════

    class Fase1Tab(Frame):
        def __init__(self, parent, **kw):
            super().__init__(parent, bg=C_PANEL, **kw)
            self._out = ""
            self._build()

        def _build(self):
            # Header tab
            hdr = Frame(self, bg=C_BG, pady=10)
            hdr.pack(fill="x")
            Label(hdr, text="FASE 1", font=("Segoe UI", 13, "bold"),
                  bg=C_BG, fg=C_GOLD).pack(side="left", padx=(20, 6))
            Label(hdr, text="Convertir documento a PDF",
                  font=("Segoe UI", 12), bg=C_BG, fg=C_TEXT).pack(side="left")

            # Separador
            Frame(self, bg=C_BORDER, height=1).pack(fill="x", padx=18)

            # Info box
            ib = Frame(self, bg=C_CARD2, padx=16, pady=10)
            ib.pack(fill="x", padx=20, pady=(14, 0))
            Label(ib, text="DOCX  →  PDF     •     PPTX  →  PDF",
                  font=("Segoe UI", 10, "bold"),
                  bg=C_CARD2, fg=C_ACCENT).pack(anchor="w")
            Label(ib,
                  text="Con LibreOffice: texto + imágenes + tipografía original\n"
                       "Sin LibreOffice: solo texto (motor Python)",
                  font=("Segoe UI", 8), bg=C_CARD2, fg=C_MUTED,
                  justify="left").pack(anchor="w", pady=(4, 0))

            # Drop zone
            self.drop = DropZone(self, [".docx", ".pptx"],
                                 "Documento a convertir")
            self.drop.pack(fill="x", padx=20, pady=16, ipady=4)

            # Botón
            btn_row = Frame(self, bg=C_PANEL)
            btn_row.pack(pady=4)
            Button(btn_row, text="  ⚡  Convertir a PDF  ",
                   font=("Segoe UI", 12, "bold"),
                   bg=C_ACCENT, fg="white", activebackground="#3A6AEE",
                   relief="flat", padx=20, pady=9,
                   cursor="hand2", command=self._run).pack()

            # Status
            self.status = StatusBar(self)
            self.status.pack(fill="x", padx=18, pady=(12, 0))

            # Resultado
            self.lbl_res = Label(self, text="",
                                  font=("Segoe UI", 9),
                                  bg=C_PANEL, fg=C_GREEN, wraplength=480)
            self.lbl_res.pack(pady=6)

            Button(self, text="📂  Abrir carpeta de destino",
                   font=("Segoe UI", 9), bg=C_CARD, fg=C_MUTED,
                   activebackground=C_BORDER,
                   relief="flat", padx=12, pady=4,
                   cursor="hand2",
                   command=lambda: _open_folder(Path(self._out))).pack()

        def _run(self):
            src = self.drop.get()
            if not src:
                messagebox.showwarning("Sin archivo", "Selecciona un archivo primero.")
                return
            dst = filedialog.asksaveasfilename(
                title="Guardar PDF como…",
                initialfile=Path(src).stem + ".pdf",
                defaultextension=".pdf",
                filetypes=[("PDF files", "*.pdf")])
            if not dst: return
            self._out = dst
            self.lbl_res.config(text="")

            def run():
                try:
                    fase1_convert(src, dst,
                        progress_cb=lambda p, m: self.after(
                            0, lambda p=p, m=m: self.status.set_status(p, m)))
                    self.after(0, lambda: self.lbl_res.config(
                        text=f"✅  {Path(dst).name}  guardado", fg=C_GREEN))
                except Exception as e:
                    log.exception("Fase 1 (escritorio) falló: %s", src)
                    self.after(0, lambda: self.lbl_res.config(
                        text=f"❌  {e}", fg=C_WARN))
                    self.after(0, lambda: self.status.set_status(0, "Error"))

            threading.Thread(target=run, daemon=True).start()


    # ══════════════════════════════════════════════════════════════════════════════
    #  TAB FASE 2
    # ══════════════════════════════════════════════════════════════════════════════

    class Fase2Tab(Frame):
        def __init__(self, parent, **kw):
            super().__init__(parent, bg=C_PANEL, **kw)
            self._out = ""
            self._build()

        def _build(self):
            # Header tab
            hdr = Frame(self, bg=C_BG, pady=10)
            hdr.pack(fill="x")
            Label(hdr, text="FASE 2", font=("Segoe UI", 13, "bold"),
                  bg=C_BG, fg=C_GOLD).pack(side="left", padx=(20, 6))
            Label(hdr, text="Descomponer PDF — PowerPoint o Word editable",
                  font=("Segoe UI", 12), bg=C_BG, fg=C_TEXT).pack(side="left")

            Frame(self, bg=C_BORDER, height=1).pack(fill="x", padx=18)

            # Info box
            ib = Frame(self, bg=C_CARD2, padx=16, pady=10)
            ib.pack(fill="x", padx=20, pady=(14, 0))
            Label(ib, text="PDF  →  PPTX  (layout visual)      PDF  →  Word  (texto corrido)",
                  font=("Segoe UI", 10, "bold"),
                  bg=C_CARD2, fg=C_ACCENT).pack(anchor="w")
            Label(ib,
                  text="• Word: texto corrido editable; Word reacomoda el flujo al editar\n"
                       "• Modo fondo fiel: imagen completa de página + texto editable\n"
                       "• Modo objetos: sin imagen base, con fotos/tablas/gráficos como objetos independientes\n"
                       "• Texto en párrafos: una caja editable por bloque, con ancho correlacionado al original\n"
                       "• Reordena, elimina o agrega páginas en PowerPoint",
                  font=("Segoe UI", 8), bg=C_CARD2, fg=C_MUTED,
                  justify="left").pack(anchor="w", pady=(4, 0))

            # Poppler status
            pcolor = C_GREEN if POPPLER_PATH else C_WARN
            pmsg   = (f"✅  Poppler detectado" if POPPLER_PATH
                      else "⚠️  Poppler no encontrado — requerido solo para modo fondo fiel")
            Label(self, text=pmsg, font=("Segoe UI", 8),
                  bg=C_PANEL, fg=pcolor).pack(anchor="w", padx=22, pady=(8, 0))

            # Opciones
            opt = Frame(self, bg=C_PANEL)
            opt.pack(fill="x", padx=22, pady=(8, 0))

            # Formato de salida.
            # PPTX: corregir presentaciones conservando el layout visual.
            # DOCX: documentos de texto para reescribir; Word reacomoda el flujo.
            fmt_row = Frame(opt, bg=C_PANEL)
            fmt_row.pack(anchor="w", pady=(0, 6))
            Label(fmt_row, text="Formato de salida:",
                  bg=C_PANEL, fg=C_TEXT, font=("Segoe UI", 9, "bold")).pack(side="left")
            self.output_format_var = StringVar(value="pptx")
            for value, label in [("pptx", "PowerPoint (layout visual) \u2713"),
                                 ("docx", "Word (texto corrido)")]:
                ttk.Radiobutton(fmt_row, text=label,
                                 variable=self.output_format_var,
                                 value=value,
                                 command=self._on_format_change).pack(side="left", padx=5)

            # DPI
            dpi_row = Frame(opt, bg=C_PANEL)
            dpi_row.pack(anchor="w")
            Label(dpi_row, text="Calidad imagen:",
                  bg=C_PANEL, fg=C_TEXT, font=("Segoe UI", 9)).pack(side="left")
            self.dpi_var = StringVar(value="150")
            for v, l in [("96","Baja"), ("150","Media ✓"), ("200","Alta"), ("300","Máxima")]:
                ttk.Radiobutton(dpi_row, text=l,
                                 variable=self.dpi_var, value=v).pack(side="left", padx=5)

            # Toggle texto
            txt_row = Frame(opt, bg=C_PANEL)
            txt_row.pack(anchor="w", pady=(6, 0))
            self.txt_var = StringVar(value="1")
            ttk.Checkbutton(txt_row,
                             text="Incluir texto editable",
                             variable=self.txt_var,
                             onvalue="1", offvalue="0").pack(side="left")

            # Modo de agrupación del texto embebido.
            # Párrafos: recomendado para edición real.
            # Líneas: conserva el comportamiento original del MVP.
            mode_row = Frame(opt, bg=C_PANEL)
            mode_row.pack(anchor="w", pady=(6, 0))
            Label(mode_row, text="Agrupar texto:",
                  bg=C_PANEL, fg=C_TEXT, font=("Segoe UI", 9)).pack(side="left")
            self.text_mode_var = StringVar(value="paragraphs")
            for value, label in [("paragraphs", "Párrafos editables ✓"),
                                 ("lines", "Líneas exactas")]:
                ttk.Radiobutton(mode_row, text=label,
                                 variable=self.text_mode_var,
                                 value=value).pack(side="left", padx=5)

            # Composición de salida.
            # Fondo fiel: comportamiento estable de v3.2.
            # Objetos: elimina la imagen base y extrae fotos/tablas JPG/gráficos raster.
            comp_row = Frame(opt, bg=C_PANEL)
            comp_row.pack(anchor="w", pady=(6, 0))
            Label(comp_row, text="Composición:",
                  bg=C_PANEL, fg=C_TEXT, font=("Segoe UI", 9)).pack(side="left")
            self.composition_var = StringVar(value="objects")
            for value, label in [("background", "Fondo fiel"),
                                 ("objects", "Objetos independientes EXP")]:
                ttk.Radiobutton(comp_row, text=label,
                                 variable=self.composition_var,
                                 value=value).pack(side="left", padx=5)

            # v3.3: conserva saltos visuales dentro de una sola caja de párrafo.
            # Mejora mucho la estabilidad de ancho/estructura sin volver a mini cajas.
            br_row = Frame(opt, bg=C_PANEL)
            br_row.pack(anchor="w", pady=(6, 0))
            self.line_breaks_var = StringVar(value="1")
            ttk.Checkbutton(br_row,
                             text="Mantener saltos visuales dentro del párrafo",
                             variable=self.line_breaks_var,
                             onvalue="1", offvalue="0").pack(side="left")

            # v3.5: opción exclusiva de la salida Word.
            # Inserta las imágenes del PDF intercaladas en el texto corrido.
            img_row = Frame(opt, bg=C_PANEL)
            img_row.pack(anchor="w", pady=(6, 0))
            self.word_images_var = StringVar(value="1")
            ttk.Checkbutton(img_row,
                             text="Incluir imágenes del PDF en el documento Word",
                             variable=self.word_images_var,
                             onvalue="1", offvalue="0").pack(side="left")

            # Filas que solo aplican a una de las dos salidas.
            self._pptx_only_rows = [txt_row, mode_row, comp_row, br_row]
            self._docx_only_rows = [img_row]
            self._on_format_change()

            # Drop zone
            self.drop = DropZone(self, [".pdf"], "PDF a descomponer")
            self.drop.pack(fill="x", padx=20, pady=12, ipady=4)

            # Botón
            btn_row = Frame(self, bg=C_PANEL)
            btn_row.pack(pady=4)
            Button(btn_row, text="  🔓  Descomponer PDF  ",
                   font=("Segoe UI", 12, "bold"),
                   bg=C_GREEN, fg="#071A12", activebackground="#00A87E",
                   relief="flat", padx=20, pady=9,
                   cursor="hand2", command=self._run).pack()

            # Status
            self.status = StatusBar(self)
            self.status.pack(fill="x", padx=18, pady=(12, 0))

            # Resultado
            self.lbl_res = Label(self, text="",
                                  font=("Segoe UI", 9),
                                  bg=C_PANEL, fg=C_GREEN, wraplength=480)
            self.lbl_res.pack(pady=6)

            Button(self, text="📂  Abrir carpeta de destino",
                   font=("Segoe UI", 9), bg=C_CARD, fg=C_MUTED,
                   activebackground=C_BORDER,
                   relief="flat", padx=12, pady=4,
                   cursor="hand2",
                   command=lambda: _open_folder(Path(self._out))).pack()

        def _set_rows_state(self, rows, enabled: bool):
            """Habilita o deshabilita todos los controles de un grupo de filas."""
            state = "normal" if enabled else "disabled"
            for row in rows:
                for child in row.winfo_children():
                    try:
                        child.configure(state=state)
                    except Exception:
                        pass

        def _on_format_change(self):
            """Muestra solo las opciones relevantes al formato de salida elegido."""
            is_word = self.output_format_var.get() == "docx"
            self._set_rows_state(self._pptx_only_rows, not is_word)
            self._set_rows_state(self._docx_only_rows, is_word)

        def _run(self):
            src = self.drop.get()
            if not src:
                messagebox.showwarning("Sin archivo", "Selecciona un PDF primero.")
                return

            # ── Salida Word: texto corrido editable ──────────────────────────
            if self.output_format_var.get() == "docx":
                self._run_word(src)
                return

            composition_mode = self.composition_var.get()
            if composition_mode == "background" and not POPPLER_PATH:
                messagebox.showerror("Poppler no encontrado",
                    "Verifica la instalación de Poppler en C:\\Poppler\\ o usa el modo Objetos independientes.")
                return

            dst = filedialog.asksaveasfilename(
                title="Guardar PPTX como…",
                initialfile=Path(src).stem + "_editable.pptx",
                defaultextension=".pptx",
                filetypes=[("PowerPoint", "*.pptx")])
            if not dst: return

            self._out               = dst
            dpi                     = int(self.dpi_var.get())
            include_text            = self.txt_var.get() == "1"
            text_mode               = self.text_mode_var.get()
            preserve_line_breaks    = self.line_breaks_var.get() == "1"
            self.lbl_res.config(text="")

            def run():
                try:
                    fase2_convert(src, dst, dpi=dpi,
                        include_text=include_text,
                        text_mode=text_mode,
                        composition_mode=composition_mode,
                        preserve_line_breaks=preserve_line_breaks,
                        progress_cb=lambda p, m: self.after(
                            0, lambda p=p, m=m: self.status.set_status(p, m)))

                    if composition_mode == "objects":
                        base = "objetos independientes"
                    else:
                        base = "fondo fiel"
                        #base = "objetos independientes"
                    if include_text and text_mode == "paragraphs":
                        txt = "párrafos editables"
                    elif include_text:
                        txt = "líneas exactas"
                    else:
                        txt = "sin texto editable"

                    self.after(0, lambda: self.lbl_res.config(
                        text=f"✅  {Path(dst).name}  ({base} + {txt})", fg=C_GREEN))
                except Exception as e:
                    log.exception("Fase 2 PPTX (escritorio) falló: %s", src)
                    self.after(0, lambda: self.lbl_res.config(
                        text=f"❌  {e}", fg=C_WARN))
                    self.after(0, lambda: self.status.set_status(0, "Error"))

            threading.Thread(target=run, daemon=True).start()

        def _run_word(self, src: str):
            """Conversión PDF → Word (texto corrido editable)."""
            dst = filedialog.asksaveasfilename(
                title="Guardar Word como…",
                initialfile=Path(src).stem + "_editable.docx",
                defaultextension=".docx",
                filetypes=[("Word", "*.docx")])
            if not dst:
                return

            self._out       = dst
            include_images  = self.word_images_var.get() == "1"
            image_dpi       = int(self.dpi_var.get())
            self.lbl_res.config(text="")

            def run():
                try:
                    fase2_word_convert(
                        src, dst,
                        include_images=include_images,
                        image_dpi=image_dpi,
                        progress_cb=lambda p, m: self.after(
                            0, lambda p=p, m=m: self.status.set_status(p, m)))

                    extra = "texto + imágenes" if include_images else "solo texto"
                    self.after(0, lambda: self.lbl_res.config(
                        text=f"✅  {Path(dst).name}  (Word editable — {extra})",
                        fg=C_GREEN))
                except Exception as e:
                    log.exception("Fase 2 Word (escritorio) falló: %s", src)
                    self.after(0, lambda: self.lbl_res.config(
                        text=f"❌  {e}", fg=C_WARN))
                    self.after(0, lambda: self.status.set_status(0, "Error"))

            threading.Thread(target=run, daemon=True).start()


    # ══════════════════════════════════════════════════════════════════════════════
    #  VENTANA PRINCIPAL
    # ══════════════════════════════════════════════════════════════════════════════

    class App(Tk):
        def __init__(self):
            super().__init__()
            self.title("PDF Quick Disassembler")
            self.geometry("660x800")
            self.minsize(600, 730)
            self.configure(bg=C_BG)
            self._styles()
            self._build()

        def _styles(self):
            s = ttk.Style(self)
            s.theme_use("clam")
            s.configure("TNotebook",     background=C_BG,    borderwidth=0)
            s.configure("TNotebook.Tab", background=C_PANEL, foreground=C_MUTED,
                         padding=[18, 9], font=("Segoe UI", 10, "bold"))
            s.map("TNotebook.Tab",
                  background=[("selected", C_CARD)],
                  foreground=[("selected", C_GOLD)])
            s.configure("Horizontal.TProgressbar",
                         troughcolor=C_BORDER, background=C_ACCENT,
                         borderwidth=0, thickness=9)
            for w in ("TRadiobutton", "TCheckbutton"):
                s.configure(w, background=C_PANEL, foreground=C_TEXT,
                             font=("Segoe UI", 9))
                s.map(w, background=[("active", C_PANEL)],
                          foreground=[("active", C_GREEN)])

        def _build(self):
            # ── Splash header ─────────────────────────────────────────────────────
            splash = Frame(self, bg=C_BG, pady=0)
            splash.pack(fill="x")

            # Banda de color superior
            band = Frame(splash, bg=C_GOLD, height=4)
            band.pack(fill="x")

            content = Frame(splash, bg=C_BG, pady=14)
            content.pack(fill="x")

            # Logo + título
            left = Frame(content, bg=C_BG)
            left.pack(side="left", padx=(20, 0))
            Label(left, text="⚡", font=("Segoe UI Emoji", 32),
                  bg=C_BG, fg=C_GOLD).pack(side="left", padx=(0, 10))

            titles = Frame(content, bg=C_BG)
            titles.pack(side="left")
            Label(titles, text="PDF Quick Disassembler",
                  font=("Segoe UI", 18, "bold"),
                  bg=C_BG, fg=C_TEXT).pack(anchor="w")
            Label(titles, text="by  GoTopo  •  v3.6 Tablas Fieles  •  2026 R.Montagne",
                  font=("Segoe UI", 9),
                  bg=C_BG, fg=C_GOLD).pack(anchor="w")
            Label(titles,
                  text="Desarma, edita y reconstruye documentos PDF",
                  font=("Segoe UI", 9),
                  bg=C_BG, fg=C_MUTED).pack(anchor="w", pady=(2, 0))

            # Línea separadora
            Frame(self, bg=C_GOLD, height=1).pack(fill="x")

            # ── Notebook ──────────────────────────────────────────────────────────
            nb = Notebook(self)
            nb.pack(fill="both", expand=True, pady=0)
            nb.add(Fase1Tab(nb),       text="   📄  FASE 1  —  a PDF   ")
            nb.add(Fase2Tab(nb),       text="   🔓  FASE 2  —  a PowerPoint   ")
            nb.add(TablaExcelTab(nb),  text="   📊  FASE 3  —  Excel Table Recovery   ")

            # ── Footer ────────────────────────────────────────────────────────────
            Frame(self, bg=C_GOLD, height=1).pack(fill="x")
            footer = Frame(self, bg=C_BG, pady=6)
            footer.pack(fill="x")
            Label(footer,
                  text="GoTopo  •  r.montagne  •  Costa Rica",
                  font=("Segoe UI", 8), bg=C_BG, fg=C_MUTED).pack()


    # ══════════════════════════════════════════════════════════════════════════════
    #  TAB — TABLA EXCEL
    # ══════════════════════════════════════════════════════════════════════════════

    class TablaExcelTab(Frame):
        """Pestaña para recuperar tablas financieras desde PDFs con texto embebido."""

        def __init__(self, parent, **kw):
            super().__init__(parent, bg=C_PANEL, **kw)
            self._out_dir: Optional[Path] = None
            self._build()

        def _build(self):
            # Header
            hdr = Frame(self, bg=C_BG, pady=10)
            hdr.pack(fill="x")
            Label(hdr, text="FASE 3", font=("Segoe UI", 13, "bold"),
                  bg=C_BG, fg=C_GOLD).pack(side="left", padx=(20, 6))
            Label(hdr, text="Excel Table Recovery — recuperar tablas desde PDF",
                  font=("Segoe UI", 12), bg=C_BG, fg=C_TEXT).pack(side="left")

            Frame(self, bg=C_BORDER, height=1).pack(fill="x", padx=18)

            # ── Advertencia crítica ───────────────────────────────────────────────
            warn_frame = Frame(self, bg="#2A1E0A", padx=16, pady=10)
            warn_frame.pack(fill="x", padx=20, pady=(14, 0))

            Label(warn_frame,
                  text="⚠️  SOLO para tablas copiadas directamente desde Excel",
                  font=("Segoe UI", 10, "bold"),
                  bg="#2A1E0A", fg=C_GOLD, justify="left").pack(anchor="w")
            Label(warn_frame,
                  text="Funciona cuando en Excel se seleccionó el rango → Copy → Paste en PowerPoint\n"
                       "→ se convirtió a PDF con LibreOffice o PDF Quick Disassembler.\n"
                       "NO funciona con tablas guardadas como PNG, JPG o captura de pantalla.",
                  font=("Segoe UI", 8), bg="#2A1E0A", fg="#C8A060",
                  justify="left").pack(anchor="w", pady=(4, 0))

            # ── Info de salida ────────────────────────────────────────────────────
            info = Frame(self, bg=C_CARD2, padx=16, pady=8)
            info.pack(fill="x", padx=20, pady=(8, 0))
            Label(info,
                  text="PDF  →  XLSX + CSV  (carpeta automática en el Desktop)",
                  font=("Segoe UI", 10, "bold"),
                  bg=C_CARD2, fg=C_ACCENT).pack(anchor="w")
            Label(info,
                  text="• Escanea todas las páginas automáticamente\n"
                       "• Genera un archivo Excel con una hoja por tabla encontrada\n"
                       "• También genera un CSV por página como respaldo",
                  font=("Segoe UI", 8), bg=C_CARD2, fg=C_MUTED,
                  justify="left").pack(anchor="w", pady=(3, 0))

            # ── Opciones avanzadas ────────────────────────────────────────────────
            adv_frame = Frame(self, bg=C_PANEL)
            adv_frame.pack(fill="x", padx=20, pady=(8, 0))

            Label(adv_frame, text="Opciones avanzadas:",
                  font=("Segoe UI", 9, "bold"),
                  bg=C_PANEL, fg=C_MUTED).pack(anchor="w")

            row1 = Frame(adv_frame, bg=C_PANEL)
            row1.pack(anchor="w", pady=(4, 0))

            # Tolerancia de gap
            Label(row1, text="Separación columnas:",
                  bg=C_PANEL, fg=C_TEXT, font=("Segoe UI", 9)).pack(side="left")
            self.gap_var = StringVar(value="18")
            ttk.Radiobutton(row1, text="Estrecha (12)",
                             variable=self.gap_var, value="12").pack(side="left", padx=4)
            ttk.Radiobutton(row1, text="Normal (18) ✓",
                             variable=self.gap_var, value="18").pack(side="left", padx=4)
            ttk.Radiobutton(row1, text="Amplia (28)",
                             variable=self.gap_var, value="28").pack(side="left", padx=4)

            row2 = Frame(adv_frame, bg=C_PANEL)
            row2.pack(anchor="w", pady=(4, 0))
            Label(row2, text="Tamaño máx. fuente tabla:",
                  bg=C_PANEL, fg=C_TEXT, font=("Segoe UI", 9)).pack(side="left")
            self.font_var = StringVar(value="16")
            for v, l in [("12", "12pt"), ("16", "16pt ✓"), ("22", "22pt")]:
                ttk.Radiobutton(row2, text=l,
                                 variable=self.font_var, value=v).pack(side="left", padx=4)

            # ── Drop zone ─────────────────────────────────────────────────────────
            self.drop = DropZone(self, [".pdf"], "PDF con tablas de Excel")
            self.drop.pack(fill="x", padx=20, pady=12, ipady=4)

            # ── Botón ─────────────────────────────────────────────────────────────
            btn_row = Frame(self, bg=C_PANEL)
            btn_row.pack(pady=4)
            Button(btn_row, text="  📊  Recuperar Tablas Excel  ",
                   font=("Segoe UI", 12, "bold"),
                   bg=C_GOLD, fg="#1A1200", activebackground="#D4A800",
                   relief="flat", padx=20, pady=9,
                   cursor="hand2", command=self._run).pack()

            # ── Status ────────────────────────────────────────────────────────────
            self.status = StatusBar(self)
            self.status.pack(fill="x", padx=18, pady=(12, 0))

            # ── Resultado ─────────────────────────────────────────────────────────
            self.lbl_res = Label(self, text="",
                                  font=("Segoe UI", 9),
                                  bg=C_PANEL, fg=C_GREEN, wraplength=500)
            self.lbl_res.pack(pady=6)

            Button(self, text="📂  Abrir carpeta de salida",
                   font=("Segoe UI", 9), bg=C_CARD, fg=C_MUTED,
                   activebackground=C_BORDER,
                   relief="flat", padx=12, pady=4,
                   cursor="hand2",
                   command=self._open_out).pack()

        def _get_desktop_output(self, pdf_stem: str) -> Path:
            """Crea carpeta de salida automáticamente en el Desktop."""
            desktop = Path.home() / "Desktop"
            if not desktop.exists():
                desktop = Path.home()
            folder = desktop / f"TablaExcel_{pdf_stem}"
            folder.mkdir(parents=True, exist_ok=True)
            return folder

        def _run(self):
            src = self.drop.get()
            if not src:
                messagebox.showwarning("Sin archivo",
                                       "Selecciona un PDF primero.")
                return

            out_dir = self._get_desktop_output(Path(src).stem)
            self._out_dir = out_dir
            self.lbl_res.config(text="")

            gap_tol      = float(self.gap_var.get())
            max_font_size = float(self.font_var.get())

            def run():
                try:
                    result = tabla_excel_recover(
                        src_path=src,
                        out_dir=out_dir,
                        gap_tol=gap_tol,
                        max_font_size=max_font_size,
                        progress_cb=lambda p, m: self.after(
                            0, lambda p=p, m=m: self.status.set_status(p, m)),
                    )
                    if result["msg"] == "sin_tablas":
                        self.after(0, lambda: self.lbl_res.config(
                            text="⚠️  No se encontraron tablas compatibles.\n"
                                 "Verifica que la tabla haya sido copiada directamente desde Excel.",
                            fg=C_WARN))
                        return

                    n     = result["recovered"]
                    xlsx  = result["xlsx"]
                    warns = result.get("warn_cells", 0)
                    bands = result.get("band_pages", 0)

                    metodo = ("rejilla del PDF" if not bands
                              else f"{result.get('grid_pages', 0)} por rejilla, "
                                   f"{bands} estimada(s)")
                    msg = (f"✅  {n} tabla(s) — {metodo}  →  "
                           f"{out_dir.name}  (Desktop)")
                    if xlsx:
                        msg += f"\n📊  {xlsx.name}"

                    # El aviso de celdas dudosas va en ámbar: un dato mal
                    # recuperado que pase inadvertido invalida la tabla entera.
                    if warns:
                        msg += (f"\n⚠️  {warns} celda(s) marcada(s) en ámbar en el "
                                f"Excel. Verifíquelas contra el PDF.")
                        color = C_GOLD
                    else:
                        msg += "\n✓  Sin celdas dudosas detectadas."
                        color = C_GREEN

                    self.after(0, lambda m=msg, c=color:
                                self.lbl_res.config(text=m, fg=c))

                except Exception as e:
                    log.exception("Table Recovery (escritorio) falló: %s", src)
                    self.after(0, lambda: self.lbl_res.config(
                        text=f"❌  Error: {e}", fg=C_WARN))
                    self.after(0, lambda: self.status.set_status(0, "Error"))

            threading.Thread(target=run, daemon=True).start()

        def _open_out(self):
            if self._out_dir and self._out_dir.exists():
                _open_folder(self._out_dir)
            else:
                folder = Path.home() / "Desktop"
                _open_folder(folder)


    # ══════════════════════════════════════════════════════════════════════════════
    def main():
        missing = []
        for mod, pkg in [("pypdf","pypdf"), ("docx","python-docx"),
                         ("pptx","python-pptx"), ("PIL","Pillow"),
                         ("pdf2image","pdf2image"), ("reportlab","reportlab"),
                         ("pdfplumber","pdfplumber"), ("fitz","PyMuPDF"),
                         ("openpyxl","openpyxl")]:
            try: __import__(mod)
            except ImportError: missing.append(pkg)
        if missing:
            log.warning("Falta instalar:  pip install %s", " ".join(missing))
        App().mainloop()

    if __name__ == "__main__":
        main()
