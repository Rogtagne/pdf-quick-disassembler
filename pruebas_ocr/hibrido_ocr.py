"""Prototipo híbrido: OCR solo donde el PDF no tiene texto digital, con dos lectores.

Por cada página:
  - Las tablas con texto digital quedan para el Table Recovery actual (exacto, sin OCR).
  - Cada imagen incrustada grande se recorta a alta resolución y se pasa por PaddleOCR-VL,
    que encuentra las tablas y su estructura.
  - Cada tabla se vuelve a leer con OCR clásico (PP-OCRv5) y se arma por coordenadas:
    cada número es una caja real en la imagen, así que no puede "repetir" valores.
  - Se elige la lectura que mejor cuadra con la columna "Total"; en el Excel se marcan en
    ámbar las celdas cuyo total no cuadra y las celdas donde los dos lectores no coinciden.

Uso (con venv_ocr):
    venv_ocr\\Scripts\\python.exe pruebas_ocr\\hibrido_ocr.py PDF [PDF ...] [--min-area 0.15] [--dpi 200]
"""
import argparse
import re
import time
from io import StringIO
from pathlib import Path

import fitz  # PyMuPDF

AMBAR = "FFC000"
NUM_RE = re.compile(r"^\(?-?\s*[\d\s.,]+\)?$")
# El OCR lee el signo menos como raya, guion largo o incluso el carácter chino "一".
GUIONES = str.maketrans({c: "-" for c in "–—−‐‑一"})


def recortar_imagenes(pdf: Path, out_dir: Path, min_area: float, dpi: int):
    """Devuelve [(pagina, n_palabras, ruta_png)] para cada imagen incrustada grande."""
    recortes = []
    doc = fitz.open(pdf)
    for i, page in enumerate(doc):
        area_pag = page.rect.width * page.rect.height
        n_palabras = len(page.get_text("words"))
        bloques = [b for b in page.get_text("dict")["blocks"] if b["type"] == 1]
        for k, b in enumerate(bloques):
            rect = fitz.Rect(b["bbox"]) & page.rect
            if rect.width * rect.height / area_pag < min_area:
                continue
            png = out_dir / f"p{i + 1:02d}_img{k + 1}.png"
            page.get_pixmap(dpi=dpi, clip=rect, alpha=False).save(png)
            recortes.append((i + 1, n_palabras, png))
    return recortes


def a_numero(txt: str):
    """'1 234,56' / '- 339 752' / '(1 234)' → float; None si no es número."""
    t = str(txt).strip().translate(GUIONES)
    if not t or not NUM_RE.match(t) or not re.search(r"\d", t):
        return None
    neg = t.startswith("-") or (t.startswith("(") and t.endswith(")"))
    t = re.sub(r"[()\-\s]", "", t)
    if "," in t and "." in t:                      # 1.234,56 o 1,234.56
        dec = "," if t.rfind(",") > t.rfind(".") else "."
        t = t.replace("." if dec == "," else ",", "").replace(dec, ".")
    elif "," in t:                                 # 1234,56 (coma decimal) o 1,234 (miles)
        t = t.replace(",", ".") if len(t.split(",")[-1]) != 3 else t.replace(",", "")
    try:
        v = float(t)
    except ValueError:
        return None
    return -v if neg else v


def tablas_html(html: str):
    import pandas as pd
    try:
        return pd.read_html(StringIO(html), header=None)
    except ValueError:
        return []


def columna_total(df):
    """Índice de la columna cuyo encabezado (primeras 3 filas) es 'Total', o None."""
    for r in range(min(3, len(df))):
        for c in range(df.shape[1] - 1, -1, -1):
            if str(df.iat[r, c]).strip().lower() == "total":
                return c
    return None


def comprobar_totales(df):
    """(filas que cuadran, celdas (fila, col) cuyo total no coincide con la suma de la fila)."""
    cuadran, dudosas = set(), []
    col_total = columna_total(df)
    if col_total is None:
        return cuadran, dudosas
    for r in range(len(df)):
        total = a_numero(df.iat[r, col_total])
        partes = [a_numero(df.iat[r, c]) for c in range(col_total)]
        partes = [p for p in partes if p is not None]
        if total is None or len(partes) < 2:
            continue
        if abs(sum(partes) - total) > max(1.0, abs(total) * 0.0005):
            dudosas.append((r, col_total))
        else:
            cuadran.add(r)
    return cuadran, dudosas


def verificar_totales(df):
    """Celdas (fila, col) dudosas: la columna 'Total' no coincide con la suma de la fila."""
    return comprobar_totales(df)[1]


def hay_guion(gris, x0, x1, y0, y1):
    """True si en la zona hay un trazo horizontal corto aislado (signo menos contable)."""
    import cv2
    import numpy as np
    if x1 - x0 < 4:
        return False
    zona = (gris[int(y0):int(y1), int(x0):int(x1)] < 128).astype(np.uint8)
    n, _, stats, _ = cv2.connectedComponentsWithStats(zona, connectivity=8)
    for x, y, ancho, alto, area in stats[1:]:
        # Aislado (no es un trozo de la regla de la tabla), a media altura y con forma de guion.
        aislado = x > 0 and x + ancho < zona.shape[1]
        centrado = 0.3 <= (y + alto / 2) / zona.shape[0] <= 0.8
        if (aislado and centrado and 5 <= ancho <= 30
                and alto <= max(4, ancho / 3) and area >= ancho * 0.8):
            return True
    return False


def tabla_por_coordenadas(textos, cajas, gris=None):
    """Arma la tabla a partir de cajas de OCR clásico [x0, y0, x1, y1].

    Filas: cajas cuyo centro vertical está a menos de media altura de línea.
    Columnas: la fila de encabezado más ancha (p. ej. meses) define los centros; cada caja
    numérica va a la columna más cercana; lo que queda a la izquierda es la etiqueta.
    Devuelve None si no hay un encabezado de al menos 4 columnas.
    """
    import numpy as np
    import pandas as pd

    textos = [str(t).translate(GUIONES) for t in textos]
    items = sorted(zip(textos, cajas), key=lambda t: (t[1][1] + t[1][3]) / 2)
    if not items:
        return None
    alto = np.median([c[3] - c[1] for _, c in items])
    filas, actual, y_ref = [], [], None
    for txt, c in items:
        yc = (c[1] + c[3]) / 2
        if y_ref is not None and yc - y_ref > alto * 0.5:
            filas.append(sorted(actual, key=lambda t: t[1][0]))
            actual = []
        actual.append((txt, c))
        y_ref = yc if len(actual) == 1 else (y_ref * (len(actual) - 1) + yc) / len(actual)
    if actual:
        filas.append(sorted(actual, key=lambda t: t[1][0]))

    def es_num(t):
        return a_numero(t) is not None or t.strip() == "-"

    candidatas = [f for f in filas if sum(not es_num(t) for t, _ in f) >= 4]
    if not candidatas:
        return None
    enc = max(candidatas, key=len)
    centros = [(c[0] + c[2]) / 2 for t, c in enc]
    paso = np.median(np.diff(centros))
    # Columnas = la racha más larga de encabezados con separación regular; lo que hay a la
    # izquierda (n.º de fila, "GASTOS GENERALES") es etiqueta.
    mejor, ini = (0, 1), 0
    for i in range(1, len(centros) + 1):
        if i == len(centros) or not (0.6 * paso <= centros[i] - centros[i - 1] <= 1.4 * paso):
            if i - ini > mejor[1] - mejor[0]:
                mejor = (ini, i)
            ini = i
    cols = centros[mejor[0]:mejor[1]]
    if len(cols) < 4:
        return None
    limite_izq = cols[0] - paso / 2

    tabla = []
    for f in filas:
        # "22 126 710 -": el guion final es una caja aparte pegada a la derecha.
        toks = []
        for t, c in f:
            if re.match(r"^.*\d\s+-$", t):
                toks += [[t[:-1].strip(), list(c)], ["-", [c[2] - 8, c[1], c[2], c[3]]]]
            else:
                toks.append([t, list(c)])
        # Formato contable: el "-" va al borde izquierdo de su celda, así que es el signo del
        # número que tiene justo a la derecha. Un "-" sin número cerca es un cero y se queda.
        for i, (t, c) in enumerate(toks):
            if t != "-" or i + 1 >= len(toks):
                continue
            sig_t, sig_c = toks[i + 1]
            if (a_numero(sig_t) is not None and not sig_t.startswith("-")
                    and sig_c[0] - c[2] < paso):
                toks[i + 1][0] = "-" + sig_t
                toks[i][0] = None
        toks = [tk for tk in toks if tk[0] is not None]
        # Respaldo por píxeles: el detector a veces no ve el guion suelto a la izquierda.
        if gris is not None:
            for i, (t, c) in enumerate(toks):
                if i and a_numero(t) is not None and not t.startswith("-") \
                        and hay_guion(gris, toks[i - 1][1][2] + 3, c[0] - 2, c[1], c[3]):
                    toks[i][0] = "-" + t

        etiqueta, celdas = [], [[] for _ in cols]
        for t, c in toks:
            xc = (c[0] + c[2]) / 2
            if xc < limite_izq:
                etiqueta.append(t)
            else:
                celdas[int(np.argmin([abs(xc - x) for x in cols]))].append(t)
        tabla.append([" ".join(etiqueta), *(" ".join(ts) for ts in celdas)])
    return pd.DataFrame(tabla)


def filas_numericas(df):
    """[(fila, [(col, valor), ...])] de las filas con al menos 2 números."""
    salida = []
    for r in range(len(df)):
        nums = [(c, a_numero(df.iat[r, c])) for c in range(df.shape[1])]
        nums = [(c, v) for c, v in nums if v is not None]
        if all(v == int(v) and 1900 <= v <= 2100 for _, v in nums):
            continue                                  # fila de años (encabezado), no datos
        if len(nums) >= 2:
            salida.append((r, nums))
    return salida


def discrepancias(final, otra):
    """Celdas de `final` cuyo valor no coincide con la fila más parecida de `otra`.

    Las dos lecturas pueden tener columnas distintas (n.º de fila aparte, etc.), así que se
    comparan las secuencias de números de cada fila, no posiciones absolutas.
    """
    otras = [[v for _, v in nums] for _, nums in filas_numericas(otra)]
    if not otras:
        return []
    dudosas = []
    for r, nums in filas_numericas(final):
        valores = [v for _, v in nums]
        if valores in otras:
            continue
        # Fila de la otra lectura con más valores iguales en la misma posición.
        par = max(otras, key=lambda o: sum(a == b for a, b in zip(valores, o)))
        if sum(a == b for a, b in zip(valores, par)) < len(valores) / 2:
            continue                                  # sin pareja clara: no se compara
        for k, (c, v) in enumerate(nums):
            if k >= len(par) or par[k] != v:
                dudosas.append((r, c))
    return dudosas


def exportar_excel(tablas, destino: Path):
    """tablas: [(pagina, png, lector, df, {(fila, col): motivo})]."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"
    resumen.append(["Hoja", "Página", "Imagen", "Lector usado", "Filas", "Columnas", "Celdas dudosas"])
    for c in resumen[1]:
        c.font = Font(bold=True)
    fill = PatternFill("solid", fgColor=AMBAR)

    for n, (pagina, png, lector, df, dudosas) in enumerate(tablas, 1):
        ws = wb.create_sheet(f"T{n:02d}_p{pagina}")
        for r in range(len(df)):
            fila = []
            for c in range(df.shape[1]):
                v = df.iat[r, c]
                v = "" if str(v) == "nan" else v
                num = a_numero(v)
                fila.append(num if num is not None and r > 0 else v)
            ws.append(fila)
        for (r, c), motivo in dudosas.items():
            celda = ws.cell(row=r + 1, column=c + 1)
            celda.fill = fill
            celda.comment = Comment(motivo, "OCR")
        resumen.append([ws.title, pagina, png.name, lector, len(df), df.shape[1], len(dudosas)])
    wb.save(destino)


def leer_tabla(ocr, imagen, bloque, html_df):
    """Combina la lectura de PaddleOCR-VL con la de PP-OCRv5 por coordenadas.

    Devuelve (lector, df, {(fila, col): motivo}).
    """
    import numpy as np
    x0, y0, x1, y1 = (int(v) for v in bloque["block_bbox"])
    recorte = imagen[max(0, y0 - 4):y1 + 4, max(0, x0 - 4):x1 + 4]
    # PaddleOCR espera BGR contiguo: con una vista recortada devuelve coordenadas erróneas.
    res = next(iter(ocr.predict(np.ascontiguousarray(recorte[:, :, ::-1]))))
    d = res.json["res"]
    coord = tabla_por_coordenadas(d["rec_texts"], d["rec_boxes"], recorte.mean(axis=2))

    # Las coordenadas ganan solo si la tabla tiene "Total" y cuadra igual o mejor: ahí no
    # repiten valores. Sin totales, el VL maneja mejor encabezados de 2 líneas y celdas unidas.
    lector, df = "PaddleOCR-VL", html_df
    if (coord is not None and columna_total(coord) is not None
            and len(verificar_totales(coord)) <= len(verificar_totales(html_df))):
        lector, df = "PP-OCRv5 coordenadas", coord

    cuadran, malas = comprobar_totales(df)
    dudosas = {}
    if coord is not None:
        otra = html_df if df is coord else coord
        for r, c in discrepancias(df, otra):
            if r not in cuadran:                      # si la fila cuadra, el total la confirma
                dudosas[(r, c)] = "Los dos lectores de OCR no coinciden en esta celda"
    for rc in malas:
        dudosas[rc] = "El total no coincide con la suma de la fila"
    return lector, df, dudosas


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdfs", nargs="+")
    ap.add_argument("--min-area", type=float, default=0.15,
                    help="Fracción mínima de la página que debe ocupar una imagen (0.15)")
    ap.add_argument("--dpi", type=int, default=200)
    args = ap.parse_args()

    import numpy as np
    from PIL import Image
    from paddleocr import PaddleOCR, PaddleOCRVL
    vl = PaddleOCRVL(device="gpu:0")
    ocr = PaddleOCR(use_doc_orientation_classify=False, use_doc_unwarping=False,
                    use_textline_orientation=False, device="gpu:0")

    for pdf in map(Path, args.pdfs):
        out = Path(__file__).parent / "resultados_hibrido" / pdf.stem.strip()
        (out / "recortes").mkdir(parents=True, exist_ok=True)
        recortes = recortar_imagenes(pdf, out / "recortes", args.min_area, args.dpi)
        print(f"\n== {pdf.name}: {len(recortes)} imágenes a leer")

        tablas = []
        t0 = time.perf_counter()
        for pagina, n_palabras, png in recortes:
            t1 = time.perf_counter()
            imagen = np.array(Image.open(png).convert("RGB"))
            for res in vl.predict(str(png)):
                for b in res.json["res"]["parsing_res_list"]:
                    if b["block_label"] != "table":
                        continue
                    dfs = tablas_html(b["block_content"])
                    if not dfs:
                        continue
                    lector, df, dudosas = leer_tabla(ocr, imagen, b, max(dfs, key=len))
                    tablas.append((pagina, png, lector, df, dudosas))
            print(f"  p{pagina:>2} {png.name}: {time.perf_counter() - t1:5.1f} s "
                  f"(texto digital en la página: {n_palabras} palabras)")

        destino = out / f"{pdf.stem.strip()}_tablas_ocr.xlsx"
        if tablas:
            exportar_excel(tablas, destino)
        total = time.perf_counter() - t0
        dud = sum(len(t[4]) for t in tablas)
        print(f"  → {len(tablas)} tablas, {dud} celdas dudosas, {total:.0f} s en total")
        for pagina, png, lector, df, dudosas in tablas:
            print(f"     p{pagina:>2} {png.name}: {lector}, {df.shape}, {len(dudosas)} dudosas")
        if tablas:
            print(f"  → {destino}")


if __name__ == "__main__":
    main()
